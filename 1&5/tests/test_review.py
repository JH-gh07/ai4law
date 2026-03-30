from io import BytesIO

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from backend.services.file_service import FileService


def build_docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("第一条 数据出境说明")
    document.add_paragraph("本协议涉及境外接收方，但未写明联系方式。")
    document.add_paragraph("第二条 安全措施")
    document.add_paragraph("双方应采取必要时的安全措施。")
    handle = BytesIO()
    document.save(handle)
    return handle.getvalue()


def build_pdf_bytes() -> bytes:
    handle = BytesIO()
    pdf = canvas.Canvas(handle, pagesize=letter)
    pdf.drawString(72, 720, "Clause 1 Data processing purpose")
    pdf.drawString(72, 700, "The processing purpose is customer service and basic information handling.")
    pdf.save()
    return handle.getvalue()


def test_review_docx_workflow(client):
    task_id = client.post("/api/v1/review/tasks").json()["id"]
    files = {"file": ("review.docx", build_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    upload_response = client.post(f"/api/v1/review/tasks/{task_id}/files", files=files)
    assert upload_response.status_code == 200

    analyze_response = client.post(f"/api/v1/review/tasks/{task_id}/analyze")
    assert analyze_response.status_code == 200
    assert analyze_response.json()["status"] == "COMPLETED"

    status_response = client.get(f"/api/v1/review/tasks/{task_id}/status")
    assert status_response.status_code == 200
    assert status_response.json()["summary"]["overall_rating"] in {"高风险", "中风险", "低风险"}

    issues_response = client.get(f"/api/v1/review/tasks/{task_id}/issues")
    assert issues_response.status_code == 200
    assert len(issues_response.json()["issues"]) >= 1

    report_response = client.get(f"/api/v1/review/tasks/{task_id}/report")
    assert report_response.status_code == 200
    assert report_response.json()["report"]["artifact_type"] == "docx"


def test_file_service_extracts_pdf_text(tmp_path):
    from backend.core.settings import Settings

    settings = Settings(storage_dir=tmp_path / "storage")
    service = FileService(settings)
    path = tmp_path / "sample.pdf"
    path.write_bytes(build_pdf_bytes())
    text = service.extract_text(path)
    assert "Data processing purpose" in text
