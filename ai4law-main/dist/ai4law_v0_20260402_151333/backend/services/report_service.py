from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from sqlalchemy.orm import Session

from backend.core.json_utils import dumps
from backend.core.settings import Settings
from backend.models.report import ReportArtifactModel
from backend.repositories.report_repository import ReportRepository
from backend.schemas.common import ReportArtifact


class ReportService:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings
        self.repository = ReportRepository()

    def create_html_report(self, db: Session, owner_type: str, owner_id: str, filename: str, html: str, preview: dict) -> ReportArtifact:
        path = self._owner_dir(owner_type, owner_id) / filename
        path.write_text(html, encoding="utf-8")
        return self._persist(db, owner_type, owner_id, "html", path, preview)

    def create_pdf_report(self, db: Session, owner_type: str, owner_id: str, filename: str, lines: list[str], preview: dict) -> ReportArtifact:
        path = self._owner_dir(owner_type, owner_id) / filename
        self._write_pdf(path, lines)
        return self._persist(db, owner_type, owner_id, "pdf", path, preview)

    def create_docx_report(self, db: Session, owner_type: str, owner_id: str, filename: str, sections: list[tuple[str, list[str]]], preview: dict) -> ReportArtifact:
        path = self._owner_dir(owner_type, owner_id) / filename
        document = Document()
        for heading, paragraphs in sections:
            document.add_heading(heading, level=1)
            for paragraph in paragraphs:
                document.add_paragraph(paragraph)
        document.save(path)
        return self._persist(db, owner_type, owner_id, "docx", path, preview)

    def get_owner_artifact(self, db: Session, owner_type: str, owner_id: str, artifact_type: str) -> ReportArtifact | None:
        model = self.repository.get_by_owner_and_type(db, owner_type, owner_id, artifact_type)
        return self._to_schema(model) if model else None

    def _owner_dir(self, owner_type: str, owner_id: str) -> Path:
        path = self.settings.report_dir / owner_type / owner_id
        path.mkdir(parents=True, exist_ok=True)
        return path

    def _write_pdf(self, path: Path, lines: list[str]) -> None:
        pdf = canvas.Canvas(str(path), pagesize=A4)
        width, height = A4
        y = height - 40
        for line in lines:
            pdf.drawString(40, y, line[:90])
            y -= 18
            if y < 60:
                pdf.showPage()
                y = height - 40
        pdf.save()

    def _persist(self, db: Session, owner_type: str, owner_id: str, artifact_type: str, path: Path, preview: dict) -> ReportArtifact:
        model = ReportArtifactModel(
            owner_type=owner_type,
            owner_id=owner_id,
            artifact_type=artifact_type,
            file_path=str(path),
            preview_json=dumps(preview),
        )
        created = self.repository.create(db, model)
        return self._to_schema(created)

    def _to_schema(self, model: ReportArtifactModel) -> ReportArtifact:
        from backend.core.json_utils import loads

        return ReportArtifact(
            id=model.id,
            owner_type=model.owner_type,
            owner_id=model.owner_id,
            artifact_type=model.artifact_type,
            file_path=model.file_path,
            preview=loads(model.preview_json, {}),
            created_at=model.created_at,
        )
