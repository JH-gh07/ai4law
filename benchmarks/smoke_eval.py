"""
本文件用于定义烟雾测试评估的核心逻辑，包括检索评估和生成评估的执行函数，以及相关的辅助函数和上下文管理器。
"""
from __future__ import annotations

import json
import os
import shutil
import tempfile
import threading
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Iterator

from docx import Document

from backend.common.rag.orchestrator import RetrievalOrchestrator
from backend.common.knowledge.v2 import RetrievalRequest
from backend.core.db import build_engine, init_db
from backend.core.resource_paths import PRODUCT_SMOKE_BENCHMARK_ROOT, PROJECT_ROOT
from backend.domains.cn.document_review.clause_reviewer import ClauseReviewer
from backend.domains.cn.document_review.rag_provider import LocalRegulationKnowledgeBase
from backend.domains.cn.security_assessment import report_renderer as assessment_report_renderer
from backend.domains.cn.security_assessment.schema import AssessmentRequest
from backend.domains.cn.security_assessment.service import AssessmentService
from backend.schemas.review import ClausePosition, ClauseType, ClassifiedClause
from benchmarks.schema import GenerationEvalCase, RetrievalEvalCase

EVAL_DIR = PRODUCT_SMOKE_BENCHMARK_ROOT
_RUNTIME_LOCK = threading.Lock()


class _DisabledLLM:
    enabled = False


class _DisabledLegalService:
    enabled = False


def _load_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"Benchmark dataset does not exist: {path}")
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                rows.append(json.loads(line))
            except json.JSONDecodeError as exc:
                raise ValueError(f"Invalid JSONL at {path}:{line_number}") from exc
    return rows


def _eval_case_paths(kind: str, target: str) -> list[Path]:
    targets = ["cn", "eu", "us"] if target == "all" else [target]
    return [EVAL_DIR / f"{kind}_eval_cases_{item}.jsonl" for item in targets]


def _load_eval_cases(kind: str, target: str) -> list[dict[str, Any]]:
    schema = RetrievalEvalCase if kind == "retrieval" else GenerationEvalCase
    rows = [schema.model_validate(row).model_dump() for path in _eval_case_paths(kind, target) for row in _load_jsonl(path)]
    if not rows:
        raise ValueError(f"No {kind} benchmark cases found for target={target}")
    case_ids = [row["case_id"] for row in rows]
    duplicates = sorted({case_id for case_id in case_ids if case_ids.count(case_id) > 1})
    if duplicates:
        raise ValueError(f"Duplicate benchmark case_id: {', '.join(duplicates)}")
    return rows


def _stage_for_case(case: dict[str, Any]) -> str:
    if case.get("task_stage"):
        return str(case["task_stage"])
    module = case["module"]
    if module.endswith("assessment") or module in {"eu_dpia", "eu_tia"}:
        return "legal_grounding"
    if "review" in module or module in {"eu_scc", "eu_bcr", "us_14117"}:
        return "clause_compare"
    return "legal_grounding"


def _recall(expected: set[str], actual: set[str]) -> float:
    return len(expected.intersection(actual)) / len(expected)


def run_retrieval_eval(target: str = "cn") -> dict[str, Any]:
    cases = _load_eval_cases("retrieval", target)
    orchestrator = RetrievalOrchestrator()
    results: list[dict[str, Any]] = []
    totals = {"recall": 0.0, "mrr": 0.0, "layer": 0, "kind": 0, "jurisdiction": 0}

    for case in cases:
        stage = _stage_for_case(case)
        module = case["module"]
        bundle = orchestrator.retrieve(
            RetrievalRequest(
                module=module,
                task_stage=stage,
                query=case["query"],
                top_k=case["top_k"],
                path=case["path"] or ("assessment" if "assessment" in module or module in {"eu_dpia", "eu_tia"} else "review"),
                document_type=case["document_type"] or ("scc_contract" if module in {"cn_review", "eu_scc"} else "other"),
                jurisdiction=case["jurisdiction"],
            )
        )
        candidates = [*bundle.standard_clauses, *bundle.legal_grounding] if stage in {"clause_compare", "issue_discovery"} else bundle.legal_grounding
        items = orchestrator.dedupe_by_source(candidates)
        retrieved_ids = [item.source_id for item in items]
        expected = set(case["must_retrieve_source_ids"])
        hit_positions = [index for index, source_id in enumerate(retrieved_ids, start=1) if source_id in expected]
        recall = _recall(expected, set(retrieved_ids))
        mrr = 1.0 / hit_positions[0] if hit_positions else 0.0

        layer_hits = [item.chunk_id for item in items if item.layer in set(case["must_not_retrieve_layers"])]
        kind_hits = [item.chunk_id for item in items if item.source_kind in set(case["must_not_retrieve_source_kinds"])]
        jurisdiction_hits = [item.chunk_id for item in items if item.jurisdiction in set(case["must_not_retrieve_jurisdictions"])]
        totals["recall"] += recall
        totals["mrr"] += mrr
        totals["layer"] += int(bool(layer_hits))
        totals["kind"] += int(bool(kind_hits))
        totals["jurisdiction"] += int(bool(jurisdiction_hits))
        results.append({
            "case_id": case["case_id"],
            "retrieved_source_ids": retrieved_ids,
            "recall_at_k": recall,
            "mrr": mrr,
            "forbidden_layer_hits": layer_hits,
            "forbidden_source_kind_hits": kind_hits,
            "cross_jurisdiction_hits": jurisdiction_hits,
        })

    count = len(cases)
    return {
        "target": target,
        "case_count": count,
        "Recall@K": round(totals["recall"] / count, 4),
        "MRR": round(totals["mrr"] / count, 4),
        "Forbidden-layer retrieval rate": round(totals["layer"] / count, 4),
        "Forbidden-source-kind retrieval rate": round(totals["kind"] / count, 4),
        "Cross-jurisdiction contamination rate": round(totals["jurisdiction"] / count, 4),
        "results": results,
    }


def _install_eval_templates() -> None:
    if assessment_report_renderer.TEMPLATE_MD.exists() and assessment_report_renderer.TEMPLATE_PATH.exists():
        return
    template_dir = Path(tempfile.gettempdir()) / "ai4law_eval_templates"
    template_dir.mkdir(parents=True, exist_ok=True)
    markdown_template = template_dir / "assessment_template.md"
    markdown_template.write_text("# {{company_name}}\n\n{{business_flow_summary}}\n\n{{overall_conclusion}}\n", encoding="utf-8")
    docx_template = template_dir / "assessment_template.docx"
    if not docx_template.exists():
        document = Document()
        document.add_heading("{{company_name}}", level=0)
        document.add_paragraph("{{business_flow_summary}}")
        document.add_paragraph("{{overall_conclusion}}")
        document.save(docx_template)
    assessment_report_renderer.TEMPLATE_MD = markdown_template
    assessment_report_renderer.TEMPLATE_PATH = docx_template


@contextmanager
def _isolated_runtime() -> Iterator[None]:
    """Keep benchmark reports, traces, and audit DB writes outside shared storage."""
    runtime_parent = PROJECT_ROOT / "outputs" / "benchmarks" / "runtime"
    runtime_parent.mkdir(parents=True, exist_ok=True)
    with _RUNTIME_LOCK, tempfile.TemporaryDirectory(prefix="run-", dir=runtime_parent) as raw_dir:
        previous_cwd = Path.cwd()
        runtime_root = Path(raw_dir)
        try:
            storage_root = runtime_root / "storage"
            storage_root.mkdir()
            shared_rag = PROJECT_ROOT / "storage" / "rag"
            if shared_rag.exists():
                shutil.copytree(shared_rag, storage_root / "rag", copy_function=shutil.copy2)
            os.chdir(runtime_root)
            engine = build_engine("sqlite:///./storage/ai4law.db")
            init_db(engine)
            engine.dispose()
            yield
        finally:
            os.chdir(previous_cwd)


def _review_input(case: dict[str, Any]) -> tuple[ClauseType, str, str]:
    payload = case["input"]
    return ClauseType(payload["clause_type"]), str(payload["clause_text"]), str(payload["document_type"])


def _issue_keys(issues: list[Any]) -> set[str]:
    return {
        value
        for issue in issues
        for value in (getattr(issue, "issue_id", ""), getattr(issue, "title", ""))
        if value
    }


def _source_layers(value: Any, result: dict[str, str] | None = None) -> dict[str, str]:
    result = result if result is not None else {}
    if isinstance(value, dict):
        source_id = value.get("source_id") or value.get("rule_id")
        layer = value.get("layer")
        if source_id and layer:
            result[str(source_id)] = str(layer)
        for child in value.values():
            _source_layers(child, result)
    elif isinstance(value, list):
        for child in value:
            _source_layers(child, result)
    return result


def _review_case_result(case: dict[str, Any], reviewer: ClauseReviewer, kb: LocalRegulationKnowledgeBase) -> dict[str, Any]:
    clause_type, clause_text, document_type = _review_input(case)
    clause = ClassifiedClause(
        clause_id=case["case_id"],
        file_id="benchmark-review",
        text=clause_text,
        clause_type=clause_type,
        position=ClausePosition(),
    )
    issues = reviewer.review(
        clause,
        use_llm=case["module"] == "cn_review",
        document_type=document_type,
        jurisdiction=case["jurisdiction"],
        module=case["module"],
    )
    lookup = kb.lookup(
        clause_type,
        clause_text,
        enrich=True,
        jurisdiction=case["jurisdiction"],
        module=case["module"],
        document_type=document_type,
    )
    citation_ids = {
        citation.source_id
        for issue in issues
        for citation in issue.structured_citations
        if citation.source_id
    }
    citation_ids.update(
        str(item["source_id"])
        for item in lookup.get("standard_clause_candidates", [])
        if item.get("source_id")
    )
    forbidden_layers = set(case["must_not_use_layers"])
    layer_by_source = _source_layers(lookup)
    forbidden_ids = sorted(source_id for source_id in citation_ids if layer_by_source.get(source_id) in forbidden_layers)
    issue_text = json.dumps([issue.model_dump(mode="json") for issue in issues], ensure_ascii=False)
    return {
        "case_id": case["case_id"],
        "issue_recall": _recall(set(case["must_find_issues"]), _issue_keys(issues)),
        "citation_correctness": _recall(set(case["must_cite_source_ids"]), citation_ids),
        "forbidden_source_leakage": float(bool(forbidden_ids)),
        "forbidden_source_ids": forbidden_ids,
        "unsupported_claim": float(any(claim in issue_text for claim in case["must_not_claim"])),
    }


def _assessment_case_result(case: dict[str, Any], service: AssessmentService) -> dict[str, Any]:
    payload = AssessmentRequest.model_validate(case["input"]["assessment_request"])
    result = service.generate_report(payload, task_id=f"benchmark-{case['case_id'].lower()}")
    issues = json.loads(Path(result.output_files["issue_list_json"]).read_text(encoding="utf-8"))
    issue_keys = {
        value
        for issue in issues
        for value in (issue.get("issue_id"), issue.get("title"))
        if value
    }
    context_pack = json.loads(Path(result.output_files["generation_basis_pack_json"]).read_text(encoding="utf-8"))
    cited_ids = {
        item["rule_id"]
        for items in context_pack.get("legal_grounding", {}).get("by_issue", {}).values()
        for item in items
        if isinstance(item, dict) and item.get("rule_id") and item.get("external_report_allowed")
    }
    forbidden_layers = set(case["must_not_use_layers"])
    forbidden_ids = sorted(
        source_id
        for source_id, layer in _source_layers(context_pack).items()
        if source_id in cited_ids and layer in forbidden_layers
    )
    sections = {
        item.get("structured_payload", {}).get("section_id") or item.get("section_id")
        for item in context_pack.get("template_context", [])
    }
    sections.discard(None)
    markdown = Path(result.output_files["markdown"]).read_text(encoding="utf-8")
    required_sections = set(case["must_use_official_template_sections"])
    return {
        "case_id": case["case_id"],
        "issue_recall": _recall(set(case["must_find_issues"]), issue_keys),
        "citation_correctness": _recall(set(case["must_cite_source_ids"]), cited_ids),
        "forbidden_source_leakage": float(bool(forbidden_ids)),
        "forbidden_source_ids": forbidden_ids,
        "template_sections_present": sorted(sections),
        "required_sections_present": sorted(required_sections.intersection(sections)),
        "unsupported_claim": float(any(claim in markdown for claim in case["must_not_claim"])),
    }


def _run_generation_eval(target: str) -> dict[str, Any]:
    cases = _load_eval_cases("generation", target)
    kb = LocalRegulationKnowledgeBase()
    reviewer = ClauseReviewer(kb)
    _install_eval_templates()
    service = AssessmentService(llm_client=_DisabledLLM(), legal_api_service=_DisabledLegalService())
    results = [
        _assessment_case_result(case, service)
        if case["module"] == "cn_assessment"
        else _review_case_result(case, reviewer, kb)
        for case in cases
    ]
    count = len(results)
    return {
        "target": target,
        "case_count": count,
        "Issue recall": round(sum(item["issue_recall"] for item in results) / count, 4),
        "Citation correctness": round(sum(item["citation_correctness"] for item in results) / count, 4),
        "Forbidden-source leakage rate": round(sum(item["forbidden_source_leakage"] for item in results) / count, 4),
        "Unsupported-claim rate": round(sum(item["unsupported_claim"] for item in results) / count, 4),
        "results": results,
    }


def run_generation_eval(target: str = "cn") -> dict[str, Any]:
    with _isolated_runtime():
        return _run_generation_eval(target)


def run_all_evals(target: str = "cn") -> dict[str, Any]:
    return {"retrieval": run_retrieval_eval(target), "generation": run_generation_eval(target)}