from pathlib import Path

from docx import Document as DocxDocument

from backend.common.render.docx_renderer import DocxRenderer
from backend.common.render.html_renderer import HtmlRenderer
from backend.common.render.markdown_renderer import MarkdownRenderer
from backend.common.render.report_model import ReportBuilder


def _report():
    return (
        ReportBuilder(title="合规报告", company_name="示例公司", module="test")
        .add_section("执行摘要")
        .add_heading("风险判断", level=3)
        .add_paragraph("存在 **高风险** <script>alert(1)</script>")
        .add_list(["第一项", "第二项"])
        .add_table(["项目", "状态"], [["传输", "高风险"]])
        .add_divider()
        .build()
    )


def test_markdown_renderer_preserves_all_block_types() -> None:
    markdown = MarkdownRenderer().render(_report())

    assert "### 风险判断" in markdown
    assert "- 第一项" in markdown
    assert "| 传输 | 高风险 |" in markdown
    assert "---" in markdown


def test_html_renderer_preserves_structure_and_escapes_untrusted_html() -> None:
    html = HtmlRenderer().render(_report())

    assert "<h3>风险判断</h3>" in html
    assert "<ul>" in html
    assert "<table>" in html
    assert "<hr />" in html
    assert "<script>" not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html


def test_docx_renderer_creates_reopenable_structured_document(tmp_path: Path) -> None:
    path = DocxRenderer().render(_report(), tmp_path / "report.docx")
    reopened = DocxDocument(path)

    assert any(paragraph.text == "风险判断" for paragraph in reopened.paragraphs)
    assert any(paragraph.style.name == "List Bullet" for paragraph in reopened.paragraphs)
    assert len(reopened.tables) == 1
    assert reopened.tables[0].cell(1, 0).text == "传输"
