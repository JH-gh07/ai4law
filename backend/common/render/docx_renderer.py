"""Render the shared report model as a DOCX document."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from backend.common.render.report_model import (
    DividerBlock,
    HeadingBlock,
    ListBlock,
    Paragraph,
    ReportDocument,
    TableBlock,
)


class DocxRenderer:
    def render(self, document: ReportDocument, output_path: Path) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        doc.add_heading(document.metadata.title, level=0)
        for section in document.sections:
            doc.add_heading(section.heading, level=min(max(section.level, 1), 9))
            for block in section.blocks:
                self._write_block(doc, block)
        doc.save(str(output_path))
        return output_path

    @staticmethod
    def _write_block(doc: Document, block) -> None:
        if isinstance(block, Paragraph):
            if block.text.strip():
                doc.add_paragraph(block.text.strip())
            return
        if isinstance(block, HeadingBlock):
            doc.add_heading(block.text, level=min(max(block.level, 1), 9))
            return
        if isinstance(block, ListBlock):
            style = "List Number" if block.ordered else "List Bullet"
            for item in block.items:
                doc.add_paragraph(item, style=style)
            return
        if isinstance(block, TableBlock):
            if not block.headers:
                return
            table = doc.add_table(rows=1 + len(block.rows), cols=len(block.headers))
            table.style = "Table Grid"
            for column, header in enumerate(block.headers):
                table.rows[0].cells[column].text = header
            for row_index, row in enumerate(block.rows, start=1):
                for column, value in enumerate(row[: len(block.headers)]):
                    table.rows[row_index].cells[column].text = value
            return
        if isinstance(block, DividerBlock):
            doc.add_paragraph("_" * 40)
            return
        raise TypeError(f"Unsupported report block: {type(block).__name__}")
