from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Iterable, Mapping

from docx import Document


def render_markdown_report(output_path: Path, title: str, sections: Iterable[tuple[str, str]]) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    lines = [f"# {title}", ""]
    for header, content in sections:
        lines.append(f"## {header}")
        lines.append(content)
        lines.append("")
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


def render_markdown_template(
    output_path: Path, template_path: Path, mapping: Mapping[str, object]
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    template = template_path.read_text(encoding="utf-8")
    normalized = {key: _normalize_value(value) for key, value in mapping.items()}
    rendered = template
    for key, value in normalized.items():
        rendered = rendered.replace(f"{{{{{key}}}}}", value)
    output_path.write_text(rendered, encoding="utf-8")
    return output_path


def render_docx_report(output_path: Path, title: str, sections: Iterable[tuple[str, str]]) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title, level=0)

    for header, content in sections:
        doc.add_heading(header, level=1)
        for line in content.splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith("- "):
                doc.add_paragraph(stripped[2:], style="List Bullet")
            else:
                doc.add_paragraph(stripped)

    doc.save(str(output_path))
    return output_path


def format_date_stamp() -> str:
    return datetime.now().strftime("%Y%m%d")


def safe_filename(value: str) -> str:
    cleaned = (value or "").strip()
    if not cleaned:
        return "report"
    cleaned = cleaned.replace("/", "_").replace("\\", "_")
    cleaned = re.sub(r"\s+", "_", cleaned)
    return cleaned


def render_docx_template(output_path: Path, template_path: Path, mapping: Mapping[str, object]) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(template_path))
    normalized = {key: _normalize_value(value) for key, value in mapping.items()}

    for paragraph in doc.paragraphs:
        _replace_placeholders_in_paragraph(paragraph, normalized)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_placeholders_in_paragraph(paragraph, normalized)

    doc.save(str(output_path))
    return output_path


def _replace_placeholders_in_paragraph(paragraph, mapping: Mapping[str, str]) -> None:
    text = paragraph.text
    if "{{" not in text:
        return
    replaced = text
    for key, value in mapping.items():
        replaced = replaced.replace(f"{{{{{key}}}}}", value)
    if replaced == text:
        return
    for run in paragraph.runs:
        run.text = ""
    parts = replaced.split("\n")
    for idx, part in enumerate(parts):
        if idx > 0:
            paragraph.add_run().add_break()
        paragraph.add_run(part)


def _normalize_value(value: object) -> str:
    if value is None:
        return ""
    return str(value)
