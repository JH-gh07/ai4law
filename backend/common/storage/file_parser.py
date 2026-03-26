from pathlib import Path

from docx import Document
from pypdf import PdfReader


class FileParser:
    SUPPORTED_SUFFIXES = {".txt", ".md", ".json", ".csv", ".pdf", ".docx"}

    def parse_text(self, file_path: str) -> str:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        suffix = path.suffix.lower()
        if suffix not in self.SUPPORTED_SUFFIXES:
            raise ValueError(
                "Unsupported file type. Supported: txt, md, json, csv, pdf, docx."
            )
        if suffix == ".pdf":
            return self._parse_pdf(path)
        if suffix == ".docx":
            return self._parse_docx(path)
        return path.read_text(encoding="utf-8", errors="ignore")

    @staticmethod
    def _parse_pdf(path: Path) -> str:
        reader = PdfReader(str(path))
        text_chunks: list[str] = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_chunks.append(page_text.strip())
        return "\n".join(text_chunks)

    @staticmethod
    def _parse_docx(path: Path) -> str:
        document = Document(str(path))
        blocks: list[str] = []

        for para in document.paragraphs:
            if para.text.strip():
                blocks.append(para.text.strip())

        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if values:
                    blocks.append(" | ".join(values))

        return "\n".join(blocks)
