from __future__ import annotations

import json
import tempfile
from pathlib import Path
from typing import Any

from backend.common.rag.orchestrator import RetrievalOrchestrator
from backend.common.knowledge.v2 import RetrievalRequest
from backend.modules.assessment.legal_grounding import build_legal_grounding
from backend.modules.assessment.schema import AssessmentRequest
from backend.modules.assessment.service import AssessmentService
from backend.modules.assessment.schema import RegulationHit
from backend.modules.assessment import report_renderer as assessment_report_renderer
from backend.common.workflow import FactItem, IssueItem
from backend.schemas.review import ClausePosition, ClassifiedClause
from backend.services.review_service.rag_provider import LocalRegulationKnowledgeBase
from backend.services.review_service.clause_reviewer import ClauseReviewer
from backend.schemas.review import ClauseType
from docx import Document

ROOT = Path(__file__).resolve().parents[3]
EVAL_DIR = ROOT / "doc" / "knowledge" / "evaluation"


class _DisabledLLM:
    enabled = False


class _DisabledLegalService:
    enabled = False


def _load_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    if not path.exists():
        return rows
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            rows.append(json.loads(line))
    return rows


def run_retrieval_eval() -> dict[str, Any]:
    cases = _load_jsonl(EVAL_DIR / "retrieval_eval_cases_cn.jsonl")
    orchestrator = RetrievalOrchestrator()
    results: list[dict[str, Any]] = []
    total_recall = 0.0
    total_mrr = 0.0
    forbidden_layer_violations = 0
    forbidden_source_kind_violations = 0
    cross_jurisdiction_violations = 0

    for case in cases:
        module = case["module"]
        stage = "legal_grounding" if module == "cn_assessment" else "clause_compare"
        bundle = orchestrator.retrieve(
            RetrievalRequest(
                module=module,
                task_stage=stage,
                query=case["query"],
                top_k=int(case.get("top_k", 5)),
                path="assessment" if module == "cn_assessment" else "review",
                document_type="scc_contract" if module == "cn_review" else "other",
            )
        )
        if module == "cn_review":
            items = orchestrator.dedupe_by_source([
                *bundle.standard_clauses,
                *bundle.legal_grounding,
            ])
        else:
            items = orchestrator.dedupe_by_source(bundle.legal_grounding)
        retrieved_ids = [item.source_id for item in items]
        expected_ids = set(case.get("must_retrieve_source_ids", []))
        hits = [idx for idx, source_id in enumerate(retrieved_ids, start=1) if source_id in expected_ids]
        recall = len(expected_ids.intersection(retrieved_ids)) / max(len(expected_ids), 1)
        mrr = 1.0 / hits[0] if hits else 0.0
        total_recall += recall
        total_mrr += mrr

        forbidden_layers = set(case.get("must_not_retrieve_layers", []))
        forbidden_source_kinds = set(case.get("must_not_retrieve_source_kinds", []))
        forbidden_jurisdictions = set(case.get("must_not_retrieve_jurisdictions", []))

        layer_hits = [item.chunk_id for item in items if item.layer in forbidden_layers]
        source_kind_hits = [item.chunk_id for item in items if item.source_kind in forbidden_source_kinds]
        jurisdiction_hits = [item.chunk_id for item in items if item.jurisdiction in forbidden_jurisdictions]
        forbidden_layer_violations += len(layer_hits)
        forbidden_source_kind_violations += len(source_kind_hits)
        cross_jurisdiction_violations += len(jurisdiction_hits)

        results.append(
            {
                "case_id": case["case_id"],
                "retrieved_source_ids": retrieved_ids,
                "recall_at_k": recall,
                "mrr": mrr,
                "forbidden_layer_hits": layer_hits,
                "forbidden_source_kind_hits": source_kind_hits,
                "cross_jurisdiction_hits": jurisdiction_hits,
            }
        )

    count = max(len(cases), 1)
    return {
        "case_count": len(cases),
        "Recall@K": round(total_recall / count, 4),
        "MRR": round(total_mrr / count, 4),
        "Forbidden-layer retrieval rate": round(forbidden_layer_violations / count, 4),
        "Forbidden-source-kind retrieval rate": round(forbidden_source_kind_violations / count, 4),
        "Cross-jurisdiction contamination rate": round(cross_jurisdiction_violations / count, 4),
        "results": results,
    }


def _assessment_payload_for_case(case_id: str) -> AssessmentRequest:
    payloads: dict[str, dict[str, Any]] = {
        "GEN-CN-ASSESS-001": {
            "company_name": "评测公司A",
            "industry": "能源",
            "is_ciio": True,
            "contains_important_data": False,
            "pii_count": 200000,
            "spi_count": 100,
            "transfer_purpose": "跨境运维支持",
            "receiver_country": "Singapore",
            "force_override_path": False,
            "uploaded_files": [],
        },
        "GEN-CN-ASSESS-002": {
            "company_name": "评测公司B",
            "industry": "SaaS",
            "is_ciio": False,
            "contains_important_data": False,
            "pii_count": 200000,
            "spi_count": 500,
            "transfer_purpose": "跨境客服",
            "receiver_country": "Singapore",
            "force_override_path": True,
            "uploaded_files": [],
        },
        "GEN-CN-ASSESS-003": {
            "company_name": "评测公司C",
            "industry": "零售",
            "is_ciio": True,
            "contains_important_data": False,
            "pii_count": 300000,
            "spi_count": 50,
            "transfer_purpose": "跨境会员运营",
            "receiver_country": "Malaysia",
            "force_override_path": False,
            "uploaded_files": [],
        },
        "GEN-CN-ASSESS-004": {
            "company_name": "评测公司D",
            "industry": "互联网",
            "is_ciio": True,
            "contains_important_data": False,
            "pii_count": 500000,
            "spi_count": 200,
            "transfer_purpose": "境外风控分析",
            "receiver_country": "Singapore",
            "force_override_path": False,
            "uploaded_files": [],
        },
        "GEN-CN-ASSESS-005": {
            "company_name": "评测公司E",
            "industry": "汽车",
            "is_ciio": False,
            "contains_important_data": True,
            "pii_count": 120000,
            "spi_count": 1000,
            "transfer_purpose": "跨境研发协作",
            "receiver_country": "Germany",
            "force_override_path": False,
            "uploaded_files": [],
        },
    }
    return AssessmentRequest.model_validate(payloads[case_id])


def _review_case_payload(case_id: str) -> tuple[ClauseType, str]:
    mapping = {
        "GEN-CN-REVIEW-001": (
            ClauseType.CROSS_BORDER_TRANSFER,
            "如本协议与主服务协议不一致，以主服务协议为准。双方可另行约定其他优先条款。"
            "境外接收方可通过其他商业安排调整标准合同义务，且无需维持与标准合同正文一致的优先适用顺序。",
        ),
        "GEN-CN-REVIEW-002": (
            ClauseType.LIABILITY,
            "乙方因个人信息出境产生的赔偿责任总额不超过年度服务费的二倍，除此之外不再承担其他责任。"
            "如发生数据主体索赔，乙方仅在前述责任上限内承担责任，并可依据商业安排主张免责。",
        ),
        "GEN-CN-REVIEW-003": (
            ClauseType.RIGHTS_REQUEST,
            "乙方收到个人信息主体的查阅、更正、删除请求后，可视情况暂缓处理，并在条件允许时再予回应。"
            "如境外接收方需内部确认，可不设固定响应时限，亦无须立即启动权利请求处理程序。",
        ),
        "GEN-CN-REVIEW-004": (
            ClauseType.RETENTION_DELETION,
            "委托处理结束后，乙方应在必要时删除相关数据；双方无需另行出具返还或删除完成证明。"
            "对于历史备份数据，可由乙方自行决定是否长期保留，且不必向甲方说明删除完成情况。",
        ),
        "GEN-CN-REVIEW-005": (
            ClauseType.CROSS_BORDER_TRANSFER,
            "乙方作为受托方应单独负责完成数据出境安全评估、标准合同签署及全部合规申报手续。"
            "个人信息处理者无需参与前述出境手续，也无需就境外接收方合规安排承担主导责任。",
        ),
    }
    return mapping[case_id]


def _install_eval_templates() -> None:
    if assessment_report_renderer.TEMPLATE_MD.exists() and assessment_report_renderer.TEMPLATE_PATH.exists():
        return
    template_dir = Path(tempfile.gettempdir()) / "ai4law_eval_templates"
    template_dir.mkdir(parents=True, exist_ok=True)

    markdown_template = template_dir / "assessment_template.md"
    if not markdown_template.exists():
        markdown_template.write_text(
            "# {{company_name}}\n\n{{business_flow_summary}}\n\n{{overall_conclusion}}\n",
            encoding="utf-8",
        )

    docx_template = template_dir / "assessment_template.docx"
    if not docx_template.exists():
        document = Document()
        document.add_heading("{{company_name}}", level=0)
        document.add_paragraph("{{business_flow_summary}}")
        document.add_paragraph("{{overall_conclusion}}")
        document.save(docx_template)

    assessment_report_renderer.TEMPLATE_MD = markdown_template
    assessment_report_renderer.TEMPLATE_PATH = docx_template


def run_generation_eval() -> dict[str, Any]:
    cases = _load_jsonl(EVAL_DIR / "generation_eval_cases_cn.jsonl")
    kb = LocalRegulationKnowledgeBase()
    reviewer = ClauseReviewer(kb)
    _install_eval_templates()
    assessment_service = AssessmentService(
        llm_client=_DisabledLLM(),
        legal_api_service=_DisabledLegalService(),
    )
    results: list[dict[str, Any]] = []
    issue_recall_total = 0.0
    citation_correct_total = 0.0
    forbidden_source_leakage_total = 0.0
    unsupported_claim_total = 0.0

    for case in cases:
        if case["module"] == "cn_assessment":
            payload = _assessment_payload_for_case(case["case_id"])
            assessment_result = assessment_service.generate_report(payload)
            issue_ids = {item["issue_id"] for item in json.loads(Path(assessment_result.output_files["issue_list_json"]).read_text(encoding="utf-8"))}
            context_pack = json.loads(Path(assessment_result.output_files["generation_basis_pack_json"]).read_text(encoding="utf-8"))
            cited_ids = {
                item["rule_id"]
                for items in context_pack.get("legal_grounding", {}).get("by_issue", {}).values()
                for item in items
                if isinstance(item, dict) and item.get("external_report_allowed")
            }
            forbidden_layers = set(case.get("must_not_use_layers", []))
            forbidden_count = sum(
                1
                for item in context_pack.get("legal_grounding_context", [])
                if item.get("layer") in forbidden_layers
            )
            sections = {item.get("structured_payload", {}).get("section_id") or item.get("section_id") for item in context_pack.get("template_context", [])}
            required_sections = set(case.get("must_use_official_template_sections", []))
            required_issues = set(case.get("must_find_issues", []))
            issue_recall = len(required_issues.intersection(issue_ids)) / max(len(required_issues), 1)
            citation_correct = len(cited_ids.intersection(case.get("must_cite_source_ids", []))) / max(len(case.get("must_cite_source_ids", [])), 1)
            markdown = Path(assessment_result.output_files["markdown"]).read_text(encoding="utf-8")
            unsupported_claim = 1.0 if any(claim in markdown for claim in case.get("must_not_claim", [])) else 0.0
            results.append(
                {
                    "case_id": case["case_id"],
                    "issue_recall": issue_recall,
                    "citation_correctness": citation_correct,
                    "forbidden_source_leakage": forbidden_count,
                    "template_sections_present": sorted(sections),
                    "required_sections_present": sorted(required_sections.intersection(sections)),
                    "unsupported_claim": unsupported_claim,
                }
            )
            issue_recall_total += issue_recall
            citation_correct_total += citation_correct
            forbidden_source_leakage_total += forbidden_count
            unsupported_claim_total += unsupported_claim
        else:
            clause_type, clause_text = _review_case_payload(case["case_id"])
            clause = ClassifiedClause(
                clause_id=case["case_id"],
                file_id="eval-review",
                text=clause_text,
                clause_type=clause_type,
                position=ClausePosition(),
            )
            document_type = "scc_contract" if case["case_id"] in {"GEN-CN-REVIEW-001", "GEN-CN-REVIEW-002", "GEN-CN-REVIEW-003"} else "dpa"
            issues = reviewer.review(clause, use_llm=True, document_type=document_type)
            issue_titles = {issue.title for issue in issues}
            structured_citations = [citation.model_dump() for issue in issues for citation in issue.structured_citations]
            lookup = kb.lookup(clause_type, clause_text, enrich=True)
            workflow_rules = lookup.get("workflow_rules", [])
            standard_candidates = lookup.get("standard_clause_candidates", [])
            required_issues = set(case.get("must_find_issues", []))
            issue_recall = len(required_issues.intersection(issue_titles)) / max(len(required_issues), 1)
            citation_correct = (
                len(
                    set(case.get("must_cite_source_ids", [])).intersection(
                        {item.get("source_id") for item in standard_candidates} |
                        {item.get("source_id") for item in structured_citations}
                    )
                )
                / max(len(case.get("must_cite_source_ids", [])), 1)
            )
            forbidden_count = sum(
                1
                for item in structured_citations
                if item.get("layer") in set(case.get("must_not_use_layers", []))
            )
            unsupported_claim = 1.0 if not issues else 0.0
            results.append(
                {
                    "case_id": case["case_id"],
                    "issue_recall": issue_recall,
                    "citation_correctness": citation_correct,
                    "forbidden_source_leakage": forbidden_count,
                    "unsupported_claim": unsupported_claim,
                }
            )
            issue_recall_total += issue_recall
            citation_correct_total += citation_correct
            forbidden_source_leakage_total += forbidden_count
            unsupported_claim_total += unsupported_claim

    count = max(len(cases), 1)
    return {
        "case_count": len(cases),
        "Issue recall": round(issue_recall_total / count, 4),
        "Citation correctness": round(citation_correct_total / count, 4),
        "Forbidden-source leakage rate": round(forbidden_source_leakage_total / count, 4),
        "Unsupported-claim rate": round(unsupported_claim_total / count, 4),
        "results": results,
    }


def run_all_evals() -> dict[str, Any]:
    retrieval = run_retrieval_eval()
    generation = run_generation_eval()
    return {"retrieval": retrieval, "generation": generation}
