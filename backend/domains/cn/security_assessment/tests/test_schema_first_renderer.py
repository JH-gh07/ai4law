"""Schema-first renderer integration — production-form acceptance test.

Exercises the full AssessmentReportRenderer pipeline with production-shaped
chapter content (``[N]`` footnotes, as produced by ``chapter_generator``),
with ``schema_first_enabled`` both on and off.  This is the Phase 2
production-form acceptance evidence required by the implementation plan.
"""

from __future__ import annotations

import json
import time
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest

from backend.common.citation.models import CitationItem
from backend.common.citation.registry import CitationRegistry as LegacyCitationRegistry
from backend.domains.cn.security_assessment.report_renderer import AssessmentReportRenderer
from backend.domains.cn.security_assessment.schema import ChapterContent, CompanyProfile, RegulationHit

# ── production-shaped fixtures ───────────────────────────────────────────────

_CID_EXPORT_5 = "CIT-CN-EXPORT-ASSESSMENT-ART5-P01"
_CID_PIPL_13 = "CIT-CN-PIPL-ART13-P01"
_CID_CYBERSEC_6 = "CIT-CN-CYBERSEC-ART6-P01"


def _prod_profile() -> CompanyProfile:
    return CompanyProfile(
        company_name="数规通科技（测试用）",
        industry="互联网",
        is_ciio=True,
        contains_important_data=True,
        pii_count=1000000,
        spi_count=10000,
        transfer_purpose="跨境客服与数据分析",
        receiver_country="Singapore",
    )


def _prod_registry() -> LegacyCitationRegistry:
    reg = LegacyCitationRegistry()
    items = [
        (_CID_EXPORT_5, "CN-REG-001", "数据出境安全评估办法", "5"),
        (_CID_PIPL_13, "CN-LAW-003", "个人信息保护法", "13"),
        (_CID_CYBERSEC_6, "CN-LAW-004", "网络安全法", "6"),
    ]
    for cid, source, title, article in items:
        reg.register(CitationItem(
            citation_id=cid, source_id=source, citation_type="law_article",
            title=title, article_no=article, authority_level="high",
            binding_force="mandatory", external_report_allowed=True,
        ))
    for cid in (_CID_EXPORT_5, _CID_PIPL_13, _CID_CYBERSEC_6):
        reg.assign_footnote_number(cid)
    return reg


def _prod_chapters() -> list[ChapterContent]:
    return [
        ChapterContent(chapter_no=1, title="出境活动概述",
                       content=("数规通科技通过跨境客服系统和数据分析平台向新加坡传输用户个人信息 [1]。"
                                "数据处理涉及普通个人信息约100万人、敏感个人信息约1万人，"
                                "属于关键信息基础设施运营者 [2] 且涉及重要数据 [3]。"),
                       citations=[_CID_EXPORT_5, _CID_PIPL_13, _CID_CYBERSEC_6], risk_level="HIGH"),
        ChapterContent(chapter_no=2, title="出境必要性与合法性基础",
                       content=("出境活动为企业核心业务必需，已获得用户明示同意 [2]。"
                                "依据《数据出境安全评估办法》第五条进行自评估 [1]。"),
                       citations=[_CID_EXPORT_5, _CID_PIPL_13], risk_level="HIGH"),
        ChapterContent(chapter_no=3, title="接收方保护能力评估",
                       content="新加坡接收方已通过ISO 27001认证，具备数据保护能力。",
                       citations=[], risk_level="MEDIUM"),
        ChapterContent(chapter_no=4, title="安全措施与传输机制",
                       content=("采用AES-256端到端加密、HSM密钥管理、访问日志审计 [3]。"
                                "已签订标准合同条款并完成备案。"),
                       citations=[_CID_CYBERSEC_6], risk_level="MEDIUM"),
        ChapterContent(chapter_no=5, title="风险识别与分析",
                       content=("主要风险包括：数据跨境传输中的政府访问风险、"
                                "接收方安全管理能力不足风险 [1] [2]。"),
                       citations=[_CID_EXPORT_5, _CID_PIPL_13], risk_level="HIGH"),
        ChapterContent(chapter_no=6, title="剩余风险与整改建议",
                       content="建议补充定期安全审计和数据泄露应急预案。",
                       citations=[], risk_level="LOW"),
        ChapterContent(chapter_no=7, title="综合评估结论",
                       content=("在补充措施有效实施的前提下，数据出境风险可控 [1] [2] [3]。"),
                       citations=[_CID_EXPORT_5, _CID_PIPL_13, _CID_CYBERSEC_6], risk_level="LOW"),
        ChapterContent(chapter_no=8, title="附件与材料清单",
                       content="数据出境合同、安全评估报告、隐私政策。",
                       citations=[], risk_level="LOW"),
    ]


REGULATIONS_FIXTURE = [
    RegulationHit(source_id="CN-REG-001", title="数据出境安全评估办法", article="第五条",
                  snippet="数据处理者向境外提供数据，应当按照本办法的规定进行安全评估。"),
    RegulationHit(source_id="CN-LAW-003", title="个人信息保护法", article="第十三条",
                  snippet="个人信息处理者应当根据个人信息的处理目的、处理方式..."),
    RegulationHit(source_id="CN-LAW-004", title="网络安全法", article="第六条",
                  snippet="国家倡导诚实守信、健康文明的网络行为..."),
]

# Minimal context_pack to satisfy build_official_report_mapping
_CONTEXT_PACK = SimpleNamespace(
    diagnosis_result={"recommended_path": "security_assessment", "risk_level": "HIGH", "rationale": "测试"},
    module_key="assessment", request_id="test", facts=[], regulations=[], issues=[],
    evidence_chain=[],
)


def _setup_templates(tmp_path, monkeypatch):
    """Install test templates so the renderer doesn't fall back to section mode."""
    md_template = tmp_path / "assessment_template.md"
    md_template.write_text(
        "# {{company_name}}\n\n{{business_flow_summary}}\n\n{{overall_conclusion}}\n\n{{citation_map}}\n",
        encoding="utf-8",
    )
    from docx import Document
    docx_template = tmp_path / "assessment_template.docx"
    doc = Document()
    doc.add_heading("{{company_name}}", level=0)
    doc.add_paragraph("{{business_flow_summary}}")
    doc.add_paragraph("{{overall_conclusion}}")
    doc.add_paragraph("{{citation_map}}")
    doc.save(docx_template)

    import backend.domains.cn.security_assessment.report_renderer as _rr_mod
    monkeypatch.setattr(_rr_mod, "TEMPLATE_MD", md_template)
    monkeypatch.setattr(_rr_mod, "TEMPLATE_PATH", docx_template)
    monkeypatch.setattr(_rr_mod, "OFFICIAL_MD_TEMPLATE", md_template)

    # Mock build_official_report_mapping to avoid context_pack requirement
    monkeypatch.setattr(
        _rr_mod, "build_official_report_mapping",
        lambda **kw: {
            "company_name": kw.get("profile").company_name,
            "overall_risk": "HIGH",
            "report_date": kw.get("date_stamp", "2026-08-08"),
            "report_id": kw.get("report_id", "test"),
        },
    )


# ── tests ────────────────────────────────────────────────────────────────────


def test_production_form_with_schema_first_generates_document_ir(tmp_path, monkeypatch):
    _setup_templates(tmp_path, monkeypatch)
    monkeypatch.chdir(tmp_path)

    registry = _prod_registry()
    profile = _prod_profile()
    chapters = _prod_chapters()

    renderer = AssessmentReportRenderer(schema_first_enabled=True, model_name="deepseek-v3")
    start_time = time.perf_counter()

    outputs = renderer.render(
        task_id="phase2-production-form-new",
        company_name=profile.company_name,
        profile=profile,
        regulations=REGULATIONS_FIXTURE,
        chapters=chapters,
        citation_registry=registry,
        context_pack=_CONTEXT_PACK,
    )

    elapsed_s = round(time.perf_counter() - start_time, 2)

    # ── document_ir.json exists and is valid ──
    ir_path = Path(outputs["document_ir_json"])
    assert ir_path.exists(), f"document_ir.json not found at {ir_path}"

    ir_data = json.loads(ir_path.read_text(encoding="utf-8"))
    assert ir_data["document_id"] == "assessment:phase2-production-form-new"
    assert ir_data["report_type"] == "assessment"
    assert ir_data["diagnostics"] == []
    assert len(ir_data["sections"]) == 8


    # ── every [N] footnote captured as a ClaimBlock citation_ref ──
    all_refs: list[str] = []
    for section in ir_data["sections"]:
        for block in section.get("blocks", []):
            if block["type"] == "claim":
                all_refs.extend(block.get("citation_refs", []))
    assert _CID_EXPORT_5 in all_refs, f"missing {_CID_EXPORT_5} ref"
    assert _CID_PIPL_13 in all_refs, f"missing {_CID_PIPL_13} ref"
    assert _CID_CYBERSEC_6 in all_refs, f"missing {_CID_CYBERSEC_6} ref"

    # ── citation_map.json matches the registry ──
    citation_map_path = Path(outputs["citation_map_json"])
    assert citation_map_path.exists()
    cmap = json.loads(citation_map_path.read_text(encoding="utf-8"))
    footnote_map = cmap.get("footnote_map", {})
    assert len(footnote_map) == 3
    assert footnote_map["1"]["citation_id"] == _CID_EXPORT_5
    assert footnote_map["2"]["citation_id"] == _CID_PIPL_13
    assert footnote_map["3"]["citation_id"] == _CID_CYBERSEC_6

    # ── user-facing artifacts ──
    for key in ("markdown", "internal_markdown", "pdf", "body_markdown", "zip"):
        assert key in outputs, f"missing output key: {key}"
        assert Path(outputs[key]).exists(), f"output file missing for {key}"

    # ── ZIP includes document_ir.json ──
    from zipfile import ZipFile
    with ZipFile(outputs["zip"]) as z:
        names = set(z.namelist())
    assert "document_ir.json" in names

    # ── manifest ──
    manifest = {
        "task_id": "phase2-production-form-new",
        "schema_first_enabled": True,
        "model": "deepseek-v3",
        "company_name": profile.company_name,
        "chapters": 8,
        "total_claim_blocks": sum(
            1 for s in ir_data["sections"] for b in s["blocks"] if b["type"] == "claim"
        ),
        "total_paragraph_blocks": sum(
            1 for s in ir_data["sections"] for b in s["blocks"] if b["type"] == "paragraph"
        ),
        "citations_registered": 3,
        "footnote_map_keys": list(footnote_map.keys()),
        "compiler_diagnostics": ir_data["diagnostics"],
        "output_keys": sorted(outputs.keys()),
        "elapsed_s": elapsed_s,
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }
    manifest_path = tmp_path / "phase2_production_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n✅ schema_first=True production-form run: {elapsed_s}s, {len(outputs)} outputs")
    print(f"   ClaimBlocks={manifest['total_claim_blocks']}, ParagraphBlocks={manifest['total_paragraph_blocks']}")


def test_production_form_without_schema_first_omits_document_ir(tmp_path, monkeypatch):
    _setup_templates(tmp_path, monkeypatch)
    monkeypatch.chdir(tmp_path)

    registry = _prod_registry()
    profile = _prod_profile()
    chapters = _prod_chapters()

    renderer = AssessmentReportRenderer(schema_first_enabled=False, model_name="legacy")
    outputs = renderer.render(
        task_id="phase2-production-form-old",
        company_name=profile.company_name,
        profile=profile,
        regulations=REGULATIONS_FIXTURE,
        chapters=chapters,
        citation_registry=registry,
        context_pack=_CONTEXT_PACK,
    )

    assert "document_ir_json" not in outputs

    from zipfile import ZipFile
    with ZipFile(outputs["zip"]) as z:
        assert "document_ir.json" not in z.namelist()

    for key in ("markdown", "internal_markdown", "pdf", "zip"):
        assert Path(outputs[key]).exists(), f"old-flow missing: {key}"

    print(f"\n✅ schema_first=False: document_ir.json correctly absent, {len(outputs)} outputs")


def test_production_form_block_count_invariant(tmp_path, monkeypatch):
    _setup_templates(tmp_path, monkeypatch)
    monkeypatch.chdir(tmp_path)
    monkeypatch.chdir(tmp_path)

    registry = _prod_registry()
    profile = _prod_profile()
    chapters = _prod_chapters()

    def _ir_block_counts():
        renderer = AssessmentReportRenderer(schema_first_enabled=True, model_name="deepseek-v3")
        outputs = renderer.render(
            task_id="phase2-invariant",
            company_name=profile.company_name,
            profile=profile,
            regulations=REGULATIONS_FIXTURE,
            chapters=chapters,
            citation_registry=registry,
            context_pack=_CONTEXT_PACK,
        )
        ir = json.loads(Path(outputs["document_ir_json"]).read_text(encoding="utf-8"))
        counts = {}
        for s in ir["sections"]:
            for b in s["blocks"]:
                counts[b["type"]] = counts.get(b["type"], 0) + 1
        return counts

    first = _ir_block_counts()
    second = _ir_block_counts()
    assert first == second, f"block count drift: {first} vs {second}"


def test_compiler_blocks_unregistered_citation_in_production_form(tmp_path, monkeypatch):
    _setup_templates(tmp_path, monkeypatch)
    monkeypatch.chdir(tmp_path)

    empty_reg = LegacyCitationRegistry()
    profile = _prod_profile()
    chapters = [
        ChapterContent(chapter_no=1, title="概述",
                       content="需要评估 [999]。",
                       citations=[], risk_level="HIGH"),
    ]

    renderer = AssessmentReportRenderer(schema_first_enabled=True, model_name="test")
    with pytest.raises(ValueError, match="CITATION_NOT_REGISTERED"):
        renderer.render(
            task_id="phase2-blocked",
            company_name=profile.company_name,
            profile=profile,
            regulations=REGULATIONS_FIXTURE,
            chapters=chapters,
            citation_registry=empty_reg,
            context_pack=_CONTEXT_PACK,
        )


def test_citation_map_footnote_count_matches_body_references(tmp_path, monkeypatch):
    _setup_templates(tmp_path, monkeypatch)
    monkeypatch.chdir(tmp_path)

    registry = _prod_registry()
    profile = _prod_profile()
    chapters = _prod_chapters()

    renderer = AssessmentReportRenderer(schema_first_enabled=True, model_name="deepseek-v3")
    outputs = renderer.render(
        task_id="phase2-footnote-count",
        company_name=profile.company_name,
        profile=profile,
        regulations=REGULATIONS_FIXTURE,
        chapters=chapters,
        citation_registry=registry,
        context_pack=_CONTEXT_PACK,
    )

    # From document_ir: all citation_refs
    ir_data = json.loads(Path(outputs["document_ir_json"]).read_text(encoding="utf-8"))
    ir_refs: set[str] = set()
    for s in ir_data["sections"]:
        for b in s["blocks"]:
            if b["type"] == "claim":
                ir_refs.update(b.get("citation_refs", []))

    # From citation_map
    cmap = json.loads(Path(outputs["citation_map_json"]).read_text(encoding="utf-8"))
    footnote_cids: set[str] = {
        item["citation_id"] for item in cmap.get("footnote_map", {}).values()
    }

    assert ir_refs == footnote_cids, (
        f"DocumentIR refs {sorted(ir_refs)} != citation_map refs {sorted(footnote_cids)}"
    )
    print(f"\n✅ Citation consistency: {len(ir_refs)} refs match across IR, citation_map")
