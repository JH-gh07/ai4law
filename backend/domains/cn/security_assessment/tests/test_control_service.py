"""task073 T08 — Assessment control 集成：control=True 产出三门并写 trace，默认关闭不变。"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from backend.common.trace.recorder import TraceRecorder
from backend.domains.cn.security_assessment import report_renderer
from backend.domains.cn.security_assessment.schema import AssessmentRequest
from backend.domains.cn.security_assessment.service import AssessmentService
from backend.domains.cn.transfer_diagnosis.service import DiagnosisService


class _DisabledLLM:
    enabled = False


class _DisabledLegalService:
    enabled = False


def _build_service() -> AssessmentService:
    return AssessmentService(
        llm_client=_DisabledLLM(), legal_api_service=_DisabledLegalService()
    )


def _disable_external_services(monkeypatch) -> None:
    monkeypatch.setattr(
        DiagnosisService, "_build_rule_explanation",
        lambda self, result, answers: result.rationale,
    )


def _install_test_templates(monkeypatch, tmp_path: Path) -> None:
    md_template = tmp_path / "assessment_template.md"
    md_template.write_text(
        "# {{company_name}}\n\n{{business_flow_summary}}\n\n{{overall_conclusion}}\n",
        encoding="utf-8",
    )
    docx_template = tmp_path / "assessment_template.docx"
    document = Document()
    document.add_heading("{{company_name}}", level=0)
    document.add_paragraph("{{business_flow_summary}}")
    document.add_paragraph("{{overall_conclusion}}")
    document.save(docx_template)
    monkeypatch.setattr(report_renderer, "TEMPLATE_MD", md_template)
    monkeypatch.setattr(report_renderer, "TEMPLATE_PATH", docx_template)


def _payload() -> AssessmentRequest:
    return AssessmentRequest(
        company_name="测试公司",
        industry="医疗科技",
        is_ciio=True,
        contains_important_data=False,
        pii_count=200000,
        spi_count=300,
        transfer_purpose="跨境客服",
        receiver_country="Singapore",
        force_override_path=False,
        uploaded_files=[],
    )


def _read_events(trace_dir: Path) -> list[dict]:
    return [
        json.loads(p.read_text(encoding="utf-8"))
        for p in sorted(trace_dir.glob("*.json"))
        if p.name != "manifest.json"
    ]


def test_control_on_produces_three_gates_and_trace(monkeypatch, tmp_path):
    _disable_external_services(monkeypatch)
    _install_test_templates(monkeypatch, tmp_path)
    service = _build_service()
    trace = TraceRecorder(tmp_path / "trace", task_id="task-c")

    result = service.generate_report(_payload(), trace=trace, control=True)

    assert result.control_decision is not None
    gates = {gr.gate for gr in result.control_decision.gate_results}
    assert gates == {"EVIDENCE_SUFFICIENCY", "CITATION_VALIDITY", "ESCALATION"}
    assert result.control_decision.legal_control_status in {
        "AUTO", "CONDITIONAL", "NEEDS_CLARIFICATION", "NEEDS_REVIEW", "BLOCKED",
    }

    events = _read_events(tmp_path / "trace")
    control_events = [
        e for e in events
        if str(e["payload"].get("detail", {}).get("raw_name", "")).startswith("control.")
    ]
    # EVIDENCE_SUFFICIENCY + CITATION_VALIDITY + ESCALATION = 3 个控制事件
    assert len(control_events) == 3
    assert {e["payload"]["detail"]["gate"] for e in control_events} == {
        "EVIDENCE_SUFFICIENCY", "CITATION_VALIDITY", "ESCALATION",
    }
    for ce in control_events:
        assert ce["name"] in {"intermediate", "warning"}


def test_control_off_has_no_control_decision(monkeypatch, tmp_path):
    _disable_external_services(monkeypatch)
    _install_test_templates(monkeypatch, tmp_path)
    service = _build_service()

    result = service.generate_report(_payload())

    assert result.control_decision is None
