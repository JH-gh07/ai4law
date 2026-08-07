import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from backend.common.citation.models import CitationItem
from backend.common.citation.registry import CitationRegistry
from backend.common.workflow.context_pack import GenerationContextPack
from backend.domains.cn.security_assessment import report_renderer
from backend.domains.cn.security_assessment.report_renderer import AssessmentReportRenderer
from backend.domains.cn.security_assessment.schema import ChapterContent, CompanyProfile, RegulationHit


def _profile() -> CompanyProfile:
    return CompanyProfile(
        company_name="测试公司",
        industry="互联网",
        is_ciio=True,
        contains_important_data=True,
        pii_count=1000000,
        spi_count=10000,
        transfer_purpose="跨境客服",
        receiver_country="Singapore",
    )


def _chapters() -> list[ChapterContent]:
    titles = [
        "出境活动概述",
        "数据类型与规模",
        "出境必要性与合法性基础",
        "境外接收方保障能力",
        "个人信息权益影响分析",
        "安全措施与传输机制",
        "剩余风险与整改建议",
        "综合评估结论",
    ]
    return [
        ChapterContent(
            chapter_no=index + 1,
            title=title,
            content=f"{title}内容 {{CIT-CN-PIPL-ART39-P1}}",
            citations=["《个人信息保护法》第39条"],
            risk_level="HIGH",
        )
        for index, title in enumerate(titles)
    ]


def _regulations() -> list[RegulationHit]:
    return [
        RegulationHit(
            source_id="CN-LAW-003",
            title="个人信息保护法",
            article="第三十九条",
            snippet="向境外提供个人信息应当取得单独同意。",
        )
    ]


def _citation_registry() -> CitationRegistry:
    registry = CitationRegistry()
    registry.register(CitationItem(
        citation_id="CIT-CN-PIPL-ART39-P1",
        source_id="CN-LAW-003",
        citation_type="law_article",
        title="个人信息保护法",
        article_no="39",
        authority_level="high",
        binding_force="mandatory",
        external_report_allowed=True,
    ))
    return registry


def _schema_chapters() -> list[ChapterContent]:
    return [ChapterContent(
        chapter_no=1,
        title="出境活动概述",
        content="企业应当开展申报前评估 {{CIT-CN-PIPL-ART39-P1}}。",
        citations=[],
        risk_level="HIGH",
    )]


def _context_pack() -> GenerationContextPack:
    return GenerationContextPack(
        module_key="cn.security_assessment",
        request_id="schema-first-test",
        facts=[],
        diagnosis_result={"recommended_path": "security_assessment", "risk_level": "HIGH"},
        regulations=[],
        issues=[],
        risk_summary={"risk_level": "HIGH"},
    )


def _install_templates(monkeypatch, tmp_path: Path) -> None:
    md_template = tmp_path / "assessment_template.md"
    md_template.write_text(
        "# {{company_name}}\n\n{{business_flow_summary}}\n\n{{overall_conclusion}}\n",
        encoding="utf-8",
    )
    docx_template = tmp_path / "assessment_template.docx"
    document = Document()
    document.add_heading("{{company_name}}", level=0)
    document.add_paragraph("{{business_flow_summary}}")
    document.add_paragraph("{{overall_conclusion}}")
    document.save(docx_template)
    monkeypatch.setattr(report_renderer, "TEMPLATE_MD", md_template)
    monkeypatch.setattr(report_renderer, "TEMPLATE_PATH", docx_template)


def test_renderer_keeps_markdown_as_external_and_avoids_duplicate_zip_names(tmp_path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    _install_templates(monkeypatch, tmp_path)
    renderer = AssessmentReportRenderer(llm_client=None)
    context_pack = GenerationContextPack(
        module_key="cn.security_assessment",
        request_id="test-task-1",
        facts=[],
        diagnosis_result={"recommended_path": "security_assessment", "risk_level": "HIGH"},
        regulations=[],
        issues=[],
        risk_summary={"risk_level": "HIGH"},
    )
    outputs = renderer.render(
        task_id="task-1",
        company_name="测试公司",
        profile=_profile(),
        regulations=_regulations(),
        chapters=_chapters(),
        path_warning=None,
        alignment_warning=None,
        issues=[],
        evidence_chain=[],
        attachment_notes=[],
        facts=[],
        diagnosis_result={"recommended_path": "security_assessment", "risk_level": "HIGH"},
        context_pack=context_pack,
    )
    assert outputs["markdown"].endswith("数据出境风险自评估报告_草案_20260603.md") or outputs["markdown"].endswith(".md")
    assert outputs["internal_markdown"].endswith("内部风险分析报告_20260603.md") or outputs["internal_markdown"].endswith(".md")
    with ZipFile(outputs["zip"]) as bundle:
        names = bundle.namelist()
    assert len(names) == len(set(names))
    assert "document_ir_json" not in outputs


def test_renderer_exposes_external_docx_and_markdown_keys(tmp_path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    _install_templates(monkeypatch, tmp_path)
    renderer = AssessmentReportRenderer(llm_client=None)
    context_pack = GenerationContextPack(
        module_key="cn.security_assessment",
        request_id="test-task-2",
        facts=[],
        diagnosis_result={"recommended_path": "scc_or_certification", "risk_level": "HIGH"},
        regulations=[],
        issues=[],
        risk_summary={"risk_level": "HIGH"},
    )
    outputs = renderer.render(
        task_id="task-2",
        company_name="测试公司",
        profile=_profile(),
        regulations=_regulations(),
        chapters=_chapters(),
        path_warning="诊断推荐路径为 scc_or_certification：请复核路径。",
        alignment_warning=None,
        issues=[],
        evidence_chain=[],
        attachment_notes=[],
        facts=[],
        diagnosis_result={"recommended_path": "scc_or_certification", "risk_level": "HIGH"},
        context_pack=context_pack,
    )
    assert "markdown" in outputs
    assert "docx" in outputs
    assert "internal_markdown" in outputs


def test_renderer_writes_document_ir_when_schema_first_is_enabled(tmp_path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    _install_templates(monkeypatch, tmp_path)
    renderer = AssessmentReportRenderer(
        llm_client=None,
        schema_first_enabled=True,
        model_name="test-model",
    )

    outputs = renderer.render(
        task_id="task-schema-first",
        company_name="测试公司",
        profile=_profile(),
        regulations=_regulations(),
        chapters=_schema_chapters(),
        citation_registry=_citation_registry(),
        context_pack=_context_pack(),
    )

    document_path = Path(outputs["document_ir_json"])
    payload = json.loads(document_path.read_text(encoding="utf-8"))
    assert payload["document_id"] == "assessment:task-schema-first"
    assert payload["diagnostics"] == []
    with ZipFile(outputs["zip"]) as bundle:
        assert "document_ir.json" in bundle.namelist()


def test_renderer_blocks_unregistered_citations_when_schema_first_is_enabled(tmp_path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    _install_templates(monkeypatch, tmp_path)
    renderer = AssessmentReportRenderer(llm_client=None, schema_first_enabled=True)

    with pytest.raises(ValueError, match="CITATION_NOT_REGISTERED"):
        renderer.render(
            task_id="task-schema-first-invalid",
            company_name="测试公司",
            profile=_profile(),
            regulations=_regulations(),
            chapters=_schema_chapters(),
            citation_registry=CitationRegistry(),
            context_pack=_context_pack(),
        )
