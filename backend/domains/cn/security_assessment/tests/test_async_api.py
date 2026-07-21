import time
from pathlib import Path

from docx import Document

from backend.domains.cn.security_assessment import router as assessment_router
from backend.domains.cn.security_assessment import report_renderer
from backend.domains.cn.security_assessment.schema import AssessmentRequest
from backend.domains.cn.security_assessment.service import AssessmentService
from backend.domains.cn.transfer_diagnosis.service import DiagnosisService


class _DisabledLLM:
    enabled = False


class _DisabledLegalService:
    enabled = False


def _disable_external_services(monkeypatch) -> None:
    monkeypatch.setattr(DiagnosisService, "_build_rule_explanation", lambda self, result, answers: result.rationale)
    assessment_router.service = AssessmentService(
        llm_client=_DisabledLLM(),
        legal_api_service=_DisabledLegalService(),
    )


def _install_test_templates(monkeypatch, tmp_path: Path) -> None:
    md_template = tmp_path / "assessment_template.md"
    md_template.write_text(
        "# {{company_name}}\n\n{{business_flow_summary}}\n\n{{overall_conclusion}}\n",
        encoding="utf-8",
    )

    docx_template = tmp_path / "assessment_template.docx"
    document = Document()
    document.add_heading("{{company_name}}", level=0)
    document.add_paragraph("{{business_flow_summary}}")
    document.add_paragraph("{{overall_conclusion}}")
    document.save(docx_template)

    monkeypatch.setattr(report_renderer, "TEMPLATE_MD", md_template)
    monkeypatch.setattr(report_renderer, "TEMPLATE_PATH", docx_template)


def test_assessment_async_flow(monkeypatch, tmp_path) -> None:
    _disable_external_services(monkeypatch)
    _install_test_templates(monkeypatch, tmp_path)
    accepted = assessment_router.service.submit_async(
        AssessmentRequest(
            company_name="AsyncCo",
            industry="SaaS",
            is_ciio=False,
            contains_important_data=False,
            pii_count=150000,
            spi_count=500,
            transfer_purpose="support",
            receiver_country="Singapore",
            force_override_path=True,
            uploaded_files=[],
        )
    )
    assert accepted.state in {"CREATED", "RUNNING"}

    last_state = None
    for _ in range(100):
        status = assessment_router.service.get_async_status(accepted.task_id)
        last_state = status.state
        if status.state == "COMPLETED":
            assert status.result is not None
            assert status.result.report_path.endswith(".docx")
            return
        if status.state == "FAILED":
            raise AssertionError(status.error)
        time.sleep(0.05)

    raise AssertionError(f"assessment async task timeout: {last_state}")
