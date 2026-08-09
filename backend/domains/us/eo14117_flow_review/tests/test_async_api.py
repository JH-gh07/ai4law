import time
from pathlib import Path

from docx import Document

from backend.domains.us.eo14117_flow_review import service as cn_flow_service
from backend.domains.us.eo14117_flow_review.schema import CNFlowRequest
from backend.domains.us.eo14117_flow_review.service import CNFlowService


class _DisabledLLM:
    enabled = False


def _install_test_templates(tmp_dir: Path, monkeypatch) -> None:
    tmp_dir.mkdir(parents=True, exist_ok=True)
    md_template = tmp_dir / "cn_flow_template.md"
    md_template.write_text("# {{business_overview}}\n\n{{risk_rating}}\n", encoding="utf-8")
    docx_template = tmp_dir / "cn_flow_template.docx"
    document = Document()
    document.add_heading("{{business_overview}}", level=1)
    document.add_paragraph("{{risk_rating}}")
    document.save(docx_template)
    monkeypatch.setattr(cn_flow_service, "TEMPLATE_MD", md_template)
    monkeypatch.setattr(cn_flow_service, "TEMPLATE_PATH", docx_template)


def test_cn_flow_async_flow(tmp_path, monkeypatch) -> None:
    _install_test_templates(tmp_path, monkeypatch)
    service = CNFlowService(llm_client=_DisabledLLM())
    accepted = service.submit_async(
        CNFlowRequest.model_validate(
            {
                "company_name": "AsyncCNFlow",
                "transfer_purpose": "全球客服与风控",
                "data_categories": ["账户信息", "设备信息"],
                "sensitive_data_flags": ["生物识别"],
                "us_person_count": 1000,
                "transaction_type": "vendor_agreement",
                "doj_data_category_by_item": {
                    "账户信息": "covered_personal_identifiers",
                    "设备信息": "not_14117_data",
                    "生物识别": "biometric_identifiers",
                },
                "recipient_entities": [
                    {
                        "entity_name": "US ServiceCo",
                        "country_region": "United States",
                        "entity_role": "processor",
                        "is_restricted_party": False,
                    }
                ],
                "transfer_chain": "CN -> US processor -> subprocessor",
                "attachments": [
                    {
                        "file_role": "data_inventory",
                        "file_name": "data.csv",
                        "file_format": "csv",
                        "storage_uri": "storage://uploads/data.csv",
                    },
                    {
                        "file_role": "entity_inventory",
                        "file_name": "entity.csv",
                        "file_format": "csv",
                        "storage_uri": "storage://uploads/entity.csv",
                    },
                ],
            }
        )
    )
    assert accepted.state in {"CREATED", "RUNNING"}

    for _ in range(200):
        status = service.get_async_status(accepted.task_id)
        if status.state == "COMPLETED":
            assert status.result is not None
            assert status.result.output_files["xlsx"].endswith(".xlsx")
            return
        if status.state == "FAILED":
            raise AssertionError(status.error)
        time.sleep(0.05)

    raise AssertionError("cn-flow async task timeout")
