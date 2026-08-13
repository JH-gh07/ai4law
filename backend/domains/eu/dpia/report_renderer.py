"""DPIA report renderer — produces DPIA draft documents, risk matrix, mitigation plan, and trace artifacts."""

from __future__ import annotations

import json
import re
from pathlib import Path
from shutil import copy2
from typing import TYPE_CHECKING
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl

from backend.common.render.report import format_date_stamp, render_docx_report, safe_filename
from backend.common.workflow.evidence import EvidenceItem
from backend.common.workflow.facts import FactItem
from backend.common.workflow.issues import IssueItem
from backend.domains.eu.dpia.schema import (
    DPIAChapterContent,
    DPIAProjectProfile,
)

if TYPE_CHECKING:
    from backend.common.citation.registry import CitationRegistry

_ISSUE_XLSX_HEADERS = [
    "问题编号", "标题", "描述", "类别", "严重程度",
    "事实引用", "规则引用", "证据引用", "建议措施", "影响输出",
]

_EVIDENCE_XLSX_HEADERS = [
    "证据编号", "主张", "事实引用", "规则引用", "结论", "置信度", "被使用于",
]

_RISK_XLSX_HEADERS = [
    "风险编号", "风险描述", "可能性", "影响程度", "风险等级", "受影响数据主体", "风险来源",
]

_MITIGATION_XLSX_HEADERS = [
    "措施编号", "描述", "目标风险", "状态", "负责方", "残余风险等级",
]

_APPENDIX_MARKDOWN_RE = re.compile(r"(?<!\w)[*_`]+|[*_`]+(?!\w)|\[\d+\]|\{\{CIT-[^}]+\}\}")


def _clean_appendix_text(value: str) -> str:
    """Strip Markdown/citation residue so an appendix value is semantic text."""
    return _APPENDIX_MARKDOWN_RE.sub("", value or "").strip()


def _format_trigger_reason(reason: object) -> str:
    if isinstance(reason, dict):
        for key in ("reason", "title", "description", "type"):
            value = str(reason.get(key) or "").strip()
            if value:
                return value
        return "未说明原因"
    return str(reason).strip() or "未说明原因"


def _render_dpia_markdown(
    output_path: Path,
    profile: DPIAProjectProfile,
    chapters: list[DPIAChapterContent],
    date_stamp: str,
    need_assessment: dict | None = None,
    citation_registry: "CitationRegistry | None" = None,
) -> str:
    """Render the DPIA draft as a Markdown document."""
    lines: list[str] = []
    lines.append(f"# {profile.project_name} — DPIA 草案 (Data Protection Impact Assessment)")
    lines.append(f"**生成日期**: {date_stamp}")
    lines.append(f"**项目目标**: {profile.project_goal}")
    lines.append("")

    if need_assessment:
        na = need_assessment
        lines.append("## 0. DPIA 必要性预判")
        lines.append(f"- DPIA 是否必需: {'是' if na.get('dpia_required') else '否'}")
        if na.get("trigger_reasons"):
            lines.append("- 触发理由:")
            for reason in na["trigger_reasons"]:
                lines.append(f"  - {_format_trigger_reason(reason)}")
        if na.get("prior_consultation_possible"):
            lines.append("- **可能需依据 GDPR Art 36 进行监管机构事先咨询**")
        if na.get("reasoning"):
            lines.append(f"\n{na['reasoning']}")
        lines.append("")

    lines.append("---")
    lines.append("")

    for chapter in chapters:
        lines.append(f"## {chapter.title}")
        lines.append("")
        lines.append(chapter.content)
        lines.append("")
        citations = chapter.citations
        if citations:
            if citation_registry is None:
                lines.append(f"*引用法规: {'; '.join(citations)}*")
            else:
                footnote_numbers = {
                    item.citation_id: number
                    for number, item in citation_registry.get_footnote_map().items()
                }
                visible_refs = [
                    f"[{footnote_numbers[citation_id]}]"
                    for citation_id in citations
                    if citation_id in footnote_numbers
                ]
                if visible_refs:
                    lines.append(f"*引用编号: {'; '.join(dict.fromkeys(visible_refs))}*")
        lines.append("")

    content = "\n".join(lines)
    output_path.write_text(content, encoding="utf-8")
    return str(output_path)


def _write_issue_list_json(issues: list[IssueItem], output_dir: Path) -> str:
    path = output_dir / "issue_list.json"
    path.write_text(
        json.dumps([issue.model_dump() for issue in issues], ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return str(path)


def _write_issue_list_xlsx(issues: list[IssueItem], output_dir: Path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet("问题清单")
    ws.title = "问题清单"
    ws.append(_ISSUE_XLSX_HEADERS)
    for issue in issues:
        ws.append([
            issue.issue_id,
            issue.title,
            issue.description,
            issue.category,
            issue.severity,
            "; ".join(issue.fact_refs),
            "; ".join(issue.rule_refs),
            "; ".join(issue.evidence_refs),
            issue.recommended_action,
            "; ".join(issue.affects_outputs),
        ])
    path = output_dir / "issue_list.xlsx"
    wb.save(path)
    return str(path)


def _write_evidence_chain_json(evidence_chain: list[EvidenceItem], output_dir: Path) -> str:
    path = output_dir / "evidence_chain.json"
    path.write_text(
        json.dumps([ev.model_dump() for ev in evidence_chain], ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return str(path)


def _write_evidence_chain_xlsx(evidence_chain: list[EvidenceItem], output_dir: Path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet("证据链")
    ws.title = "证据链"
    ws.append(_EVIDENCE_XLSX_HEADERS)
    for ev in evidence_chain:
        ws.append([
            ev.evidence_id,
            ev.claim,
            "; ".join(ev.fact_refs),
            "; ".join(ev.rule_refs),
            ev.conclusion,
            ev.confidence,
            "; ".join(ev.used_by),
        ])
    path = output_dir / "evidence_chain.xlsx"
    wb.save(path)
    return str(path)


def _write_risk_matrix_json(risk_matrix: list[dict], output_dir: Path) -> str:
    path = output_dir / "risk_matrix.json"
    path.write_text(json.dumps(risk_matrix, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def _write_risk_matrix_xlsx(risk_matrix: list[dict], output_dir: Path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet("风险矩阵")
    ws.title = "风险矩阵"
    ws.append(_RISK_XLSX_HEADERS)
    for item in risk_matrix:
        ws.append([
            item.get("risk_id", ""),
            item.get("risk_description", ""),
            item.get("likelihood", ""),
            item.get("impact", ""),
            item.get("risk_level", ""),
            item.get("affected_data_subjects", ""),
            item.get("risk_source", ""),
        ])
    path = output_dir / "risk_matrix.xlsx"
    wb.save(path)
    return str(path)


def _write_mitigation_plan_json(mitigation_plan: list[dict], output_dir: Path) -> str:
    path = output_dir / "mitigation_plan.json"
    path.write_text(json.dumps(mitigation_plan, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def _write_mitigation_plan_xlsx(mitigation_plan: list[dict], output_dir: Path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet("缓解措施计划")
    ws.title = "缓解措施计划"
    ws.append(_MITIGATION_XLSX_HEADERS)
    for item in mitigation_plan:
        ws.append([
            item.get("mitigation_id", ""),
            item.get("description", ""),
            "; ".join(item.get("target_risk_ids", [])),
            item.get("status", ""),
            item.get("responsible_party", ""),
            item.get("residual_risk_level", ""),
        ])
    path = output_dir / "mitigation_plan.xlsx"
    wb.save(path)
    return str(path)


def _write_facts_json(facts: list[FactItem], output_dir: Path) -> str:
    path = output_dir / "facts.json"
    path.write_text(
        json.dumps([fact.model_dump() for fact in facts], ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return str(path)


def _write_need_assessment_json(need_assessment: dict | None, output_dir: Path) -> str:
    path = output_dir / "dpia_need_assessment.json"
    path.write_text(
        json.dumps(need_assessment or {}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return str(path)


def _write_legal_grounding_json(legal_grounding: dict, output_dir: Path) -> str:
    path = output_dir / "legal_grounding.json"
    path.write_text(json.dumps(legal_grounding, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def _write_writing_strategy_json(writing_strategy: dict, output_dir: Path) -> str:
    path = output_dir / "writing_strategy.json"
    path.write_text(json.dumps(writing_strategy, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def _write_generation_basis_pack_json(generation_basis_pack: dict, output_dir: Path) -> str:
    path = output_dir / "generation_basis_pack.json"
    path.write_text(json.dumps(generation_basis_pack, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def _write_citation_map_json(
    citation_registry: "CitationRegistry", output_dir: Path, task_id: str = ""
) -> str:
    from backend.common.citation.output import write_citation_map_json

    footnote_map = citation_registry.get_footnote_map()
    return write_citation_map_json(
        output_dir=output_dir,
        module="dpia",
        task_id=task_id,
        footnote_map={str(num): item.to_dict() for num, item in footnote_map.items()},
        all_items=citation_registry.to_list(),
    )


def _write_document_ir_json(document_ir, output_dir: Path) -> str:
    """Persist the typed DocumentIR as an internal (non-user-facing) artifact."""
    path = output_dir / "document_ir.json"
    path.write_text(
        json.dumps(document_ir.model_dump(mode="json"), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return str(path)


class DPIAReportRenderer:
    """Render DPIA outputs: draft MD, risk matrix, mitigation plan, and full ZIP bundle."""

    def __init__(
        self,
        *,
        schema_first_enabled: bool = False,
        model_name: str = "legacy-dpia",
    ) -> None:
        self.schema_first_enabled = schema_first_enabled
        self.model_name = model_name

    @staticmethod
    def _build_input_appendix(
        profile: DPIAProjectProfile,
        generation_basis_pack: dict | None = None,
    ) -> list[tuple[str, str]]:
        """Sanitized input appendix entries for the schema-first report body.

        Surfaces the deterministic-chapter business structure (profile fields,
        risk matrix, mitigation plan, attachment summaries) as key-value
        entries. Raw request JSON / storage URIs / nested structures must never
        reach report body blocks (task068 T09).
        """
        entries: list[tuple[str, str]] = [
            ("项目名称", _clean_appendix_text(profile.project_name or "未提供")),
            ("项目目标", _clean_appendix_text(profile.project_goal or "未提供")),
        ]
        if profile.data_categories:
            entries.append(("数据类别", _clean_appendix_text("、".join(map(str, profile.data_categories)))))
        if profile.special_category_types:
            entries.append(("特殊类别数据", _clean_appendix_text("、".join(map(str, profile.special_category_types)))))
        if profile.data_subject_categories:
            entries.append(("数据主体类别", _clean_appendix_text("、".join(map(str, profile.data_subject_categories)))))
        if profile.lawful_basis:
            entries.append(("法律基础", _clean_appendix_text("、".join(map(str, profile.lawful_basis)))))
        if profile.cross_border_transfer:
            dest = _clean_appendix_text(profile.transfer_destination or "未提供目的地")
            entries.append(("跨境传输", f"是（{dest}）"))

        flags = [
            ("自动化决策", profile.automated_decision_making),
            ("系统性监控", profile.systematic_monitoring),
            ("大规模处理", profile.large_scale_processing),
            ("数据比对", profile.data_matching),
            ("新技术", profile.new_technology),
            ("弱势数据主体", profile.vulnerable_data_subjects),
        ]
        enabled = [label for label, flag in flags if flag]
        if enabled:
            entries.append(("高风险触发项", _clean_appendix_text("、".join(enabled))))
        if profile.dpo_name:
            entries.append(("DPO 姓名", _clean_appendix_text(profile.dpo_name)))

        basis = generation_basis_pack or {}
        risk_matrix = basis.get("risk_matrix") or []
        for index, item in enumerate(risk_matrix, start=1):
            if not isinstance(item, dict):
                continue
            name = item.get("risk_name") or item.get("description") or item.get("risk_id", "")
            level = item.get("overall_level") or item.get("risk_level", "未评估")
            entries.append((f"风险项 {index}", _clean_appendix_text(f"{name}（等级：{level}）")))

        mitigation_plan = basis.get("mitigation_plan") or []
        for index, entry in enumerate(mitigation_plan, start=1):
            if not isinstance(entry, dict):
                continue
            for measure in (entry.get("measures") or [])[:2]:
                if not isinstance(measure, dict):
                    continue
                text = measure.get("measure", "未提供")
                status = measure.get("status", "未提供")
                entries.append((f"缓解措施 {index}", _clean_appendix_text(f"{text}（状态：{status}）")))

        attachments = basis.get("attachment_summaries") or profile.attachment_metadata or []
        for index, item in enumerate(attachments, start=1):
            if isinstance(item, dict):
                name = item.get("source_ref") or item.get("filename") or item.get("file_name") or ""
                entries.append((f"附件 {index}", _clean_appendix_text(str(name))))
            elif str(item).strip():
                entries.append((f"附件 {index}", _clean_appendix_text(str(item))))

        return entries

    def render(
        self,
        task_id: str,
        profile: DPIAProjectProfile,
        chapters: list[DPIAChapterContent],
        issues: list[IssueItem] | None = None,
        evidence_chain: list[EvidenceItem] | None = None,
        facts: list[FactItem] | None = None,
        need_assessment: dict | None = None,
        risk_matrix: list[dict] | None = None,
        mitigation_plan: list[dict] | None = None,
        legal_grounding: dict | None = None,
        writing_strategy: dict | None = None,
        generation_basis_pack: dict | None = None,
        citation_registry: "CitationRegistry | None" = None,
        trace_manifest_path: str | None = None,
        consistency_issues: list[str] | None = None,
    ) -> dict[str, str]:
        output_dir = Path("outputs/dpia") / task_id / "outputs"
        output_dir.mkdir(parents=True, exist_ok=True)

        document_ir = None
        reporting_registry = None
        if self.schema_first_enabled:
            from backend.common.citation.registry import CitationRegistry as LegacyCitationRegistry
            from backend.common.reporting import DocumentCompiler
            from backend.domains.eu.dpia.schema_first import build_dpia_document_ir

            input_appendix = self._build_input_appendix(profile, generation_basis_pack)
            document_ir, reporting_registry = build_dpia_document_ir(
                task_id=task_id,
                project_name=profile.project_name,
                chapters=chapters,
                citation_registry=citation_registry or LegacyCitationRegistry(),
                model=self.model_name,
                input_appendix=input_appendix,
            )
            compile_result = DocumentCompiler().compile(document_ir, reporting_registry)
            if compile_result.status != "success":
                codes = ", ".join(item.code for item in compile_result.diagnostics)
                raise ValueError(f"Schema-first compiler blocked DPIA output: {codes}")

        date_stamp = format_date_stamp()
        safe_name = safe_filename(profile.project_name)
        result: dict[str, str] = {}
        if document_ir is not None:
            result["document_ir_json"] = _write_document_ir_json(document_ir, output_dir)

        # 1. Markdown, DOCX, and PDF drafts
        md_path = output_dir / f"{safe_name}_DPIA草案_{date_stamp}.md"
        pdf_path = output_dir / f"{safe_name}_DPIA草案_{date_stamp}.pdf"
        docx_path = output_dir / f"{safe_name}_DPIA草案_{date_stamp}.docx"

        if document_ir is not None and reporting_registry is not None:
            # Schema-first: all three formats render directly from the canonical
            # DocumentIR, never from legacy Markdown/template strings (task068 T09).
            from backend.common.reporting.renderers import render_docx, render_pdf
            from backend.common.reporting.renderers.markdown import MarkdownRenderer

            md_path.write_text(
                MarkdownRenderer().render(document_ir, reporting_registry),
                encoding="utf-8",
            )
            result["markdown"] = str(md_path)
            pdf_path.write_bytes(render_pdf(document_ir, reporting_registry))
            result["pdf"] = str(pdf_path)
            docx_path.write_bytes(render_docx(document_ir, reporting_registry))
            result["docx"] = str(docx_path)
        else:
            result["markdown"] = _render_dpia_markdown(
                md_path,
                profile,
                chapters,
                date_stamp,
                need_assessment,
                citation_registry,
            )

            from backend.common.render.pdf_renderer import get_pdf_renderer

            get_pdf_renderer().from_markdown(
                md_path.read_text(encoding="utf-8"),
                pdf_path,
                f"{profile.project_name} — DPIA 草案",
            )
            result["pdf"] = str(pdf_path)

            docx_sections = [
                ("项目概况", f"项目目标：{profile.project_goal}"),
                *[(chapter.title, chapter.content) for chapter in chapters],
            ]
            result["docx"] = str(
                render_docx_report(
                    docx_path,
                    f"{profile.project_name} — DPIA 草案",
                    docx_sections,
                )
            )
        # 2. Issue list
        if issues:
            result["issue_list_json"] = _write_issue_list_json(issues, output_dir)
            result["issue_list_xlsx"] = _write_issue_list_xlsx(issues, output_dir)

        # 3. Evidence chain
        if evidence_chain:
            result["evidence_chain_json"] = _write_evidence_chain_json(evidence_chain, output_dir)
            result["evidence_chain_xlsx"] = _write_evidence_chain_xlsx(evidence_chain, output_dir)

        # 4. Risk matrix
        if risk_matrix:
            result["risk_matrix_json"] = _write_risk_matrix_json(risk_matrix, output_dir)
            result["risk_matrix_xlsx"] = _write_risk_matrix_xlsx(risk_matrix, output_dir)

        # 5. Mitigation plan
        if mitigation_plan:
            result["mitigation_plan_json"] = _write_mitigation_plan_json(mitigation_plan, output_dir)
            result["mitigation_plan_xlsx"] = _write_mitigation_plan_xlsx(mitigation_plan, output_dir)

        # 6. Facts
        if facts:
            result["facts_json"] = _write_facts_json(facts, output_dir)

        # 7. Need assessment
        result["need_assessment_json"] = _write_need_assessment_json(need_assessment, output_dir)

        # 8. Legal grounding
        if legal_grounding:
            result["legal_grounding_json"] = _write_legal_grounding_json(legal_grounding, output_dir)

        # 9. Writing strategy
        if writing_strategy:
            result["writing_strategy_json"] = _write_writing_strategy_json(writing_strategy, output_dir)

        # 10. Generation basis pack
        if generation_basis_pack:
            result["generation_basis_pack_json"] = _write_generation_basis_pack_json(
                generation_basis_pack, output_dir
            )

        # 11. Citation map
        if citation_registry is not None:
            result["citation_map_json"] = _write_citation_map_json(
                citation_registry, output_dir, task_id
            )

        # 12. Trace manifest
        if trace_manifest_path:
            dest = output_dir / "trace_manifest.json"
            copy2(trace_manifest_path, dest)
            result["trace_manifest"] = str(dest)

        # 13. Consistency issues report
        if consistency_issues:
            ci_path = output_dir / "consistency_issues.json"
            ci_path.write_text(
                json.dumps(consistency_issues, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            result["consistency_issues_json"] = str(ci_path)

        # 14. ZIP bundle
        zip_path = output_dir / f"{safe_name}_DPIA输出包_{date_stamp}.zip"
        with ZipFile(zip_path, mode="w", compression=ZIP_DEFLATED) as bundle:
            seen: set[str] = set()
            for key, path_str in result.items():
                if key == "zip":
                    continue
                p = Path(path_str)
                if p.exists():
                    arcname = p.name
                    if arcname in seen:
                        continue
                    seen.add(arcname)
                    bundle.write(p, arcname=arcname)
        result["zip"] = str(zip_path)

        return result
