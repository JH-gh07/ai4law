"""Test EU SCC compliance review — three core test scenarios."""

import json
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

from backend.common.citation.models import CitationItem
from backend.common.citation.registry import CitationRegistry
from backend.domains.eu.scc_review.schema import SCCChapter
from backend.domains.eu.scc_review.schema import SCCReviewRequest
from backend.domains.eu.scc_review.service import EU_SCCService, _build_template_mapping


class _DisabledLLM:
    enabled = False


def test_scc_template_summary_uses_shared_citation_registry() -> None:
    registry = CitationRegistry()
    registry.register(
        CitationItem(
            citation_id="CIT-EU-GDPR-ART46-P01",
            source_id="gdpr",
            jurisdiction="eu",
            display_label="GDPR 第46条",
            title="GDPR",
            article_no="46",
        )
    )
    chapter = SCCChapter(
        chapter_no=2,
        title="总体合规评级",
        content="跨境传输应当采取适当保障。",
        citations=["CIT-EU-GDPR-ART46-P01"],
        risk_level="LOW",
    )
    payload = SimpleNamespace(company_name="引用测试公司")
    rule_result = SimpleNamespace(
        overall_rating="LOW",
        module_validation=SimpleNamespace(actual_module="Module One", is_correct=True),
        all_findings=[],
        clause_comparison=SimpleNamespace(deviations_found=0),
        tia_review=SimpleNamespace(has_third_country_transfer=True, tia_present=True),
    )

    mapping = _build_template_mapping(
        payload,
        [chapter],
        rule_result,
        "20260808",
        registry,
    )

    assert mapping["executive_summary"].endswith("[1]")


def test_scc_schema_first_service_exposes_document_ir(tmp_path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    service = EU_SCCService(llm_client=_DisabledLLM())
    service.schema_first_enabled = True
    payload = _make_request(
        _SCC_C2P_INDIA_HEALTH,
        declared_module="Module Two",
        exporter_role="controller",
        importer_role="processor",
        has_tia=False,
        has_supplementary=False,
    )

    result = service.generate_report(payload, task_id="schema-first-scc")

    assert "document_ir_json" in result.output_files
    assert Path(result.output_files["document_ir_json"]).exists()
    docx_path = Path(result.output_files["docx"])
    assert docx_path.stat().st_size > 0
    with ZipFile(docx_path) as bundle:
        assert "word/document.xml" in bundle.namelist()


def test_uploaded_scc_document_drives_core_review(tmp_path, monkeypatch) -> None:
    from docx import Document

    source_docx = tmp_path / "india_health_scc.docx"
    document = Document()
    for block in _SCC_C2P_INDIA_HEALTH.strip().split("\n\n"):
        document.add_paragraph(block)
    document.save(source_docx)

    monkeypatch.chdir(tmp_path)
    service = EU_SCCService(llm_client=_DisabledLLM())
    payload = _make_request(
        "Module Two placeholder without the uploaded contract clauses.",
        declared_module="Module Two",
        exporter_role="controller",
        importer_role="processor",
        has_tia=False,
        has_supplementary=False,
    )
    payload.uploaded_files = [str(source_docx)]

    result = service.generate_report(payload, task_id="uploaded-scc-document")

    finding_text = "\n".join(
        f"{finding.location} {finding.risk_analysis}" for finding in result.findings
    )
    assert "India" in finding_text
    assert "Clause 15" in finding_text
    assert result.attachment_notes
    assert "annotated_docx" in result.output_files

    citation_map = json.loads(
        Path(result.output_files["citation_map_json"]).read_text(encoding="utf-8")
    )
    cited_items = list(citation_map["footnote_map"].values())
    cited_labels = [item["display_label"] for item in cited_items]
    assert cited_items
    assert any("2021/914" in label for label in cited_labels)
    assert any("EDPB Recommendations 01/2020" in label for label in cited_labels)
    assert all("段落" not in label for label in cited_labels)
    assert all(item["can_jump"] is True for item in cited_items)


def test_source_case_packet_preserves_contract_sections_and_destination() -> None:
    from backend.common.storage.file_parser import FileParser
    from backend.domains.eu.scc_review.scc_parser import (
        _build_transfer_chain,
        parse_scc_document,
    )
    from backend.domains.eu.scc_review.scc_rule_engine import review_annexes

    repo_root = Path(__file__).resolve().parents[5]
    source_docx = (
        repo_root
        / "backend/tests/eu_scc/fixtures/case_02_india_health_scc_input.docx"
    )
    text = FileParser().parse_text(source_docx)

    document = parse_scc_document(
        text,
        declared_module="Module Two",
        exporter_role="controller",
        importer_role="processor",
    )
    transfer_chain = _build_transfer_chain(
        "controller", "processor", document
    )

    clauses = {clause.clause_no: clause.content for clause in document.clauses}
    assert {9, 14, 15}.issubset(clauses)
    assert "as soon as legally permissible" in clauses[15]
    assert "Patient unique study identifier" in document.annex_i_b.data_categories
    assert any("India" in location for location in transfer_chain.storage_locations)
    annex_findings = review_annexes(document, "Module Two").findings
    assert any(
        finding.issue_type == "special_category_misclassified"
        for finding in annex_findings
    )


def test_post_pipeline_agents_remain_inside_run_trace(tmp_path, monkeypatch) -> None:
    from backend.common.trace.context import current_trace
    from backend.common.trace.recorder import TraceRecorder

    monkeypatch.chdir(tmp_path)
    recorder = TraceRecorder(tmp_path / "trace", task_id="post-agent-trace")
    service = EU_SCCService(llm_client=_DisabledLLM())
    observed: list[bool] = []

    class _TraceCheckingRemediationAgent:
        def run(self, **_kwargs):
            observed.append(current_trace.get() is recorder)
            return {"patched_findings": [], "action_plan": [], "p0_count": 0}

    service.agents["remediation"] = _TraceCheckingRemediationAgent()
    payload = _make_request(
        _SCC_C2P_INDIA_HEALTH,
        declared_module="Module Two",
        exporter_role="controller",
        importer_role="processor",
        has_tia=False,
        has_supplementary=False,
    )

    service.generate_report(payload, task_id="post-agent-trace", trace=recorder)

    assert observed == [True]


def test_agent_json_parse_fallback_is_written_to_trace(tmp_path) -> None:
    from backend.common.runtime.run_manifest import summarize_trace
    from backend.common.trace.context import current_trace
    from backend.common.trace.recorder import TraceRecorder
    from backend.domains.eu.scc_review.agents import SCCAgentBase

    class _MalformedJSONLLM:
        enabled = True

        def chat(self, **_kwargs):
            return '{"enhanced": [{"finding_id": "broken"}'

    class _TestAgent(SCCAgentBase):
        agent_name = "trace_test_agent"

    recorder = TraceRecorder(tmp_path / "trace", task_id="agent-fallback")
    token = current_trace.set(recorder)
    try:
        result = _TestAgent(_MalformedJSONLLM())._call_llm("test")
    finally:
        current_trace.reset(token)

    assert result is None
    fallback = json.loads(
        (tmp_path / "trace" / "001_agent_fallback.json").read_text(
            encoding="utf-8"
        )
    )
    assert fallback["payload"]["detail"]["agent"] == "trace_test_agent"
    assert fallback["payload"]["detail"]["fallback"] is True
    assert summarize_trace(recorder)["fallback_count"] == 1


# ═══════════════════════════════════════════════════════════════════════
# Test data: EU 2021/914 style SCC document snippets
# ═══════════════════════════════════════════════════════════════════════

_SCC_C2C_US_AWS = """
MODULE ONE

Transfer from controller to controller

SECTION I

Clause 1 — Purpose and scope
The purpose of these standard contractual clauses is to ensure compliance with the requirements of Regulation (EU) 2016/679.

Clause 2 — Effect and invariability of the Clauses
These Clauses set out appropriate safeguards, including enforceable data subject rights.

Clause 7 — Docking clause
Additional parties may accede to these Clauses throughout the life-cycle of the contract.

Clause 9 — Use of sub-processors
The data importer shall not sub-contract without the data exporter's prior specific written authorization.

Clause 14 — Local laws and practices affecting compliance with the Clauses
The Parties warrant that they have no reason to believe that the laws and practices in the third country prevent the data importer from fulfilling its obligations.

Clause 15 — Obligations in case of access by public authorities
The data importer shall promptly notify the data exporter if it receives a legally binding request from a public authority to disclose personal data.

Clause 16 — Non-compliance with the Clauses and termination
Following termination, the data importer shall delete all personal data or return them.

Clause 17 — Governing law
These Clauses shall be governed by the law of one of the EU Member States.

Clause 18 — Choice of forum and jurisdiction
Any dispute shall be resolved by the courts of an EU Member State.

ANNEX I

A. LIST OF PARTIES
Data exporter: French Retail Group SA (Controller)
Data importer: UK Store Ltd (Controller)

B. DESCRIPTION OF TRANSFER
Data subjects: online shoppers
Data categories: order history data, contact details
Purpose: order processing and customer service
Frequency: continuous
Retention: 5 years

ANNEX II
TECHNICAL AND ORGANISATIONAL MEASURES
- TLS encryption in transit
- AES-256 encryption at rest
- Access controls with RBAC

ANNEX III
LIST OF SUB-PROCESSORS
- Amazon Web Services (AWS) — cloud hosting — United States
"""

_SCC_C2P_INDIA_HEALTH = """
MODULE TWO

Transfer from controller to processor

SECTION I

Clause 1 — Purpose and scope
The purpose of these standard contractual clauses is to ensure compliance.

Clause 7 — Docking clause
Additional parties may accede to these Clauses.

Clause 9 — Use of sub-processors
The data importer shall not sub-contract without the data exporter's prior specific written authorization.

Clause 14 — Local laws and practices
The Parties warrant that they have no reason to believe that the laws in the third country prevent the data importer from fulfilling its obligations. The data importer shall provide the assessment documentation upon request.

Clause 15 — Obligations in case of access by public authorities
The data importer shall notify the data exporter as soon as legally permissible if it receives a request from a public authority. The data importer shall review the legality at its discretion. The data importer shall provide information to the extent reasonably practicable.

Clause 16 — Non-compliance and termination
Following termination, the data importer shall delete or return data.

Clause 17 — Governing law
These Clauses shall be governed by the law of Ireland.

Clause 18 — Choice of forum and jurisdiction
Any dispute shall be resolved by the courts of Ireland.

ANNEX I

A. LIST OF PARTIES
Data exporter: German Health Research GmbH (Controller)
Data importer: India Data Analytics Pvt Ltd (Processor)

B. DESCRIPTION OF TRANSFER
Data subjects: clinical trial participants
Data categories: patient health data, clinical trial results, diagnostic records, medical history
Purpose: medical research analytics
Frequency: continuous
Retention: 10 years

ANNEX II
TECHNICAL AND ORGANISATIONAL MEASURES
- TLS encryption
- Access controls
- Data pseudonymization

ANNEX III
LIST OF SUB-PROCESSORS
- None
"""

_SCC_P2P_MULTI_PARTY_WRONG_MODULE = """
MODULE TWO

Transfer from controller to processor

SECTION I

Clause 7 — Docking clause
Additional parties may accede. Current list of additional parties:
- Nordic Retail Group (See MSA for details)

Clause 9 — Use of sub-processors
Approved sub-processors listed in Annex III.

ANNEX I

A. LIST OF PARTIES
Data exporter: Dutch Logistics BV (Processor)
Data importer: Orange Cloud Services SARL (Processor)
  Note: Orange Cloud provides cloud infrastructure to Nordic Retail Group's supply chain

B. DESCRIPTION OF TRANSFER
Data categories: logistics tracking data, warehouse inventory
Purpose: supply chain optimization

ANNEX II
TECHNICAL AND ORGANISATIONAL MEASURES
- TLS encryption
- Access controls

ANNEX III
LIST OF SUB-PROCESSORS
- Serbia Data Center doo — data hosting — Serbia
"""


# ═══════════════════════════════════════════════════════════════════════
# Test helpers
# ═══════════════════════════════════════════════════════════════════════

def _make_request(scc_text, declared_module="", exporter_role="", importer_role="", has_tia=False, has_supplementary=False):
    return SCCReviewRequest.model_validate({
        "project_name": "EU SCC Test",
        "scc_text": scc_text,
        "declared_module_type": declared_module,
        "exporter_role": exporter_role,
        "importer_role": importer_role,
        "has_tia": has_tia,
        "has_supplementary_measures": has_supplementary,
        "uploaded_files": [],
        "company_name": "Test Company",
    })


# ═══════════════════════════════════════════════════════════════════════
# Test Case 1: C2C + UK + US AWS → HIGH (TIA missing, insufficient measures)
# ═══════════════════════════════════════════════════════════════════════

def test_case1_c2c_us_aws_tia_missing():
    """C2C to UK, but data stored at US AWS — TIA missing, Schrems II measures absent."""
    payload = _make_request(
        _SCC_C2C_US_AWS,
        declared_module="Module One",
        exporter_role="controller",
        importer_role="controller",
        has_tia=False,
        has_supplementary=False,
    )
    service = EU_SCCService(llm_client=_DisabledLLM())
    result = service.generate_report(payload)

    # Module validation should pass (C2C is Module One)
    assert result.module_validation.is_correct is True

    # Overall rating must be HIGH (third country + no TIA + no supplementary measures)
    assert result.overall_rating == "HIGH"

    # Must have findings about TIA missing
    tia_findings = [f for f in result.findings if "tia_missing" in f.issue_type or "supplementary" in f.issue_type]
    assert len(tia_findings) > 0

    # Must detect US as a third country issue
    us_findings = [f for f in result.findings if "US" in f.risk_analysis or "CLOUD" in f.risk_analysis or "AWS" in f.location]
    assert len(us_findings) > 0

    # Must have sub-processor findings (AWS in US)
    sp_findings = [f for f in result.findings if "sub_processor" in f.issue_type]
    assert len(sp_findings) > 0

    onward_disclosure = [
        f for f in result.findings
        if f.location == "Annex I.B"
        and f.issue_type == "annex_incomplete"
        and "United States" in f.risk_analysis
    ]
    assert len(onward_disclosure) > 0

    # Verify chapters
    assert len(result.chapters) == 4
    assert result.report_path


# ═══════════════════════════════════════════════════════════════════════
# Test Case 2: C2P + India + Health Data + Clause 15 weakened → HIGH
# ═══════════════════════════════════════════════════════════════════════

def test_case2_c2p_india_clause15_weakened():
    """C2P to India, health data misclassified, Clause 15 weakened."""
    payload = _make_request(
        _SCC_C2P_INDIA_HEALTH,
        declared_module="Module Two",
        exporter_role="controller",
        importer_role="processor",
        has_tia=False,
        has_supplementary=False,
    )
    service = EU_SCCService(llm_client=_DisabledLLM())
    result = service.generate_report(payload)

    # Module validation should pass (C2P is Module Two)
    assert result.module_validation.is_correct is True

    # Must be HIGH due to Clause 15 weakening + health data + India third country
    assert result.overall_rating == "HIGH"

    # Must detect Clause 15 deviations
    # Check from findings
    clause_findings = [f for f in result.findings if "clause_weakened" in f.issue_type or "Clause 15" in f.location]
    assert len(clause_findings) > 0

    # Must detect clause deviation findings (Clause 15 weakening)
    clause15_findings = [f for f in result.findings if "clause_weakened" in f.issue_type or "Clause 15" in f.location]
    assert len(clause15_findings) > 0

    # India is a third country
    tia_findings = [f for f in result.findings if "tia_missing" in f.issue_type or "India" in f.risk_analysis]
    assert len(tia_findings) > 0

    assert len(result.chapters) == 4


# ═══════════════════════════════════════════════════════════════════════
# Test Case 3: P2P + Multi-party + Wrong Module → HIGH
# ═══════════════════════════════════════════════════════════════════════

def test_case3_p2p_wrong_module():
    """P2P using Module Two (wrong) + See MSA incomplete + Serbia sub-processor."""
    payload = _make_request(
        _SCC_P2P_MULTI_PARTY_WRONG_MODULE,
        declared_module="Module Two",
        exporter_role="processor",
        importer_role="processor",
        has_tia=False,
        has_supplementary=False,
    )
    service = EU_SCCService(llm_client=_DisabledLLM())
    result = service.generate_report(payload)

    # Module is WRONG: P2P should be Module Three, not Module Two
    assert result.module_validation.is_correct is False
    assert result.module_validation.expected_module == "Module Three"
    assert result.module_validation.actual_module == "Module Two"

    # Must be HIGH
    assert result.overall_rating == "HIGH"

    # Must have module mismatch finding
    module_findings = [f for f in result.findings if f.issue_type == "module_mismatch"]
    assert len(module_findings) > 0

    # Must detect "See MSA" incomplete info
    incomplete_findings = [f for f in result.findings if "incomplete" in f.issue_type or "MSA" in f.risk_analysis]
    assert len(incomplete_findings) > 0

    # Serbia sub-processor
    sub_proc_findings = [f for f in result.findings if "sub_processor" in f.issue_type or "Serbia" in f.risk_analysis]
    assert len(sub_proc_findings) > 0

    assert len(result.chapters) == 4


# ═══════════════════════════════════════════════════════════════════════
# Additional unit tests
# ═══════════════════════════════════════════════════════════════════════

def test_parser_extracts_module_type():
    """Verify parser correctly identifies Module One."""
    from backend.domains.eu.scc_review.scc_parser import parse_scc_document
    doc = parse_scc_document(_SCC_C2C_US_AWS)
    assert doc.module_type == "Module One"


def test_parser_extracts_clauses():
    """Verify parser extracts at least some clauses."""
    from backend.domains.eu.scc_review.scc_parser import parse_scc_document
    doc = parse_scc_document(_SCC_C2C_US_AWS)
    assert len(doc.clauses) > 3  # should extract at least 4 clauses


def test_parser_extracts_sub_processors():
    """Verify parser extracts AWS from Annex III."""
    from backend.domains.eu.scc_review.scc_parser import parse_scc_document
    doc = parse_scc_document(_SCC_C2C_US_AWS)
    sp_names = [sp.get("name", "").lower() for sp in doc.annex_iii.sub_processors]
    assert any("aws" in n or "amazon" in n for n in sp_names)


def test_module_validation_controller_controller():
    """C2C → Module One."""
    from backend.domains.eu.scc_review.scc_rule_engine import validate_module_selection
    from backend.domains.eu.scc_review.schema import SCCTransferChain
    chain = SCCTransferChain(exporter_role="controller", importer_role="controller")
    result = validate_module_selection(chain, "Module One")
    assert result.is_correct is True
    assert result.expected_module == "Module One"


def test_module_validation_controller_processor():
    """C2P → Module Two."""
    from backend.domains.eu.scc_review.scc_rule_engine import validate_module_selection
    from backend.domains.eu.scc_review.schema import SCCTransferChain
    chain = SCCTransferChain(exporter_role="controller", importer_role="processor")
    result = validate_module_selection(chain, "Module Two")
    assert result.is_correct is True


def test_module_validation_mismatch():
    """Processor×Processor declared as Module Two → should be Module Three."""
    from backend.domains.eu.scc_review.scc_rule_engine import validate_module_selection
    from backend.domains.eu.scc_review.schema import SCCTransferChain
    chain = SCCTransferChain(exporter_role="processor", importer_role="processor")
    result = validate_module_selection(chain, "Module Two")
    assert result.is_correct is False
    assert result.expected_module == "Module Three"


def test_clause_comparison_detects_weakening():
    """Clause 15 with 'as soon as legally permissible' should be detected."""
    from backend.domains.eu.scc_review.scc_rule_engine import compare_standard_clauses
    from backend.domains.eu.scc_review.schema import SCCDocument, SCCClause

    doc = SCCDocument(
        clauses=[
            SCCClause(clause_no=15, title="Clause 15", content="The data importer shall notify the data exporter as soon as legally permissible if it receives a request."),
        ]
    )
    result = compare_standard_clauses(doc)
    assert result.deviations_found > 0


def test_output_files_generated():
    """Verify all expected output files are produced."""
    payload = _make_request(_SCC_C2C_US_AWS, "Module One", "controller", "controller")
    service = EU_SCCService(llm_client=_DisabledLLM())
    result = service.generate_report(payload)
    assert "markdown" in result.output_files
    assert "findings_json" in result.output_files
    assert result.report_path.endswith(".md")
    from pathlib import Path

    from pypdf import PdfReader

    pdf_path = Path(result.output_files["pdf"])
    assert pdf_path.read_bytes().startswith(b"%PDF")
    assert len(PdfReader(pdf_path).pages) >= 1
