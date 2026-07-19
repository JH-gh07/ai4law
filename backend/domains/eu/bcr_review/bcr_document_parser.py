"""BCRDocumentParser — structured BCR document parsing."""

from __future__ import annotations

import re
from pathlib import Path
from uuid import uuid4

from docx import Document
from pypdf import PdfReader


class BCRStructuredChapter:
    def __init__(self, chapter_no: int = 0, title: str = "", content: str = "",
                 clauses: list[dict] | None = None) -> None:
        self.chapter_no = chapter_no
        self.title = title
        self.content = content
        self.clauses = clauses or []

    def to_dict(self) -> dict:
        return {
            "chapter_no": self.chapter_no, "title": self.title,
            "content": self.content, "clauses": self.clauses,
        }


class BCRStructuredDocument:
    def __init__(self, file_id: str = "", filename: str = "",
                 plain_text: str = "") -> None:
        self.file_id = file_id
        self.filename = filename
        self.plain_text = plain_text
        self.title: str = ""
        self.version: str = ""
        self.effective_date: str = ""
        self.chapters: list[BCRStructuredChapter] = []
        self.appendix_titles: list[str] = []

    def to_dict(self) -> dict:
        return {
            "file_id": self.file_id, "filename": self.filename,
            "title": self.title, "version": self.version,
            "effective_date": self.effective_date,
            "chapters": [c.to_dict() for c in self.chapters],
            "appendix_titles": self.appendix_titles,
        }


class BCRDocumentParser:
    _section_pattern = re.compile(
        r"(?:(?:Section|Chapter|Article|SECTION|CHAPTER|ARTICLE)\s+[IVXLCDM\d]+[\.:\)]|\d+\.\s+[A-Z][a-z])",
    )

    _known_titles = {
        "definitions": "definitions", "scope": "scope", "member": "scope",
        "binding": "binding nature", "third party": "third party beneficiary",
        "beneficiary": "third party beneficiary", "liable": "eu liable entity",
        "liability": "liability and compensation", "compensation": "liability and compensation",
        "data protection principle": "data protection principles",
        "data subject right": "data subject rights", "rights": "data subject rights",
        "complaint": "complaint handling", "cooperation": "cooperation",
        "onward": "onward transfer", "transfer": "onward transfer",
        "tia": "third country law assessment", "third country": "third country law assessment",
        "local law": "third country law assessment",
        "audit": "audit and training", "training": "audit and training",
        "update": "update mechanism", "amendment": "update mechanism",
        "termination": "termination", "exit": "termination",
        "government": "government access", "breach": "personal data breach",
        "security": "security measures",
    }

    def parse(self, file_id: str, file_path: Path) -> BCRStructuredDocument:
        ext = file_path.suffix.lower()
        if ext == ".docx":
            text = self._extract_docx(file_path)
        elif ext == ".pdf":
            text = self._extract_pdf(file_path)
        else:
            text = file_path.read_text(encoding="utf-8", errors="ignore")

        doc = BCRStructuredDocument(file_id=file_id, filename=file_path.name, plain_text=text)
        self._extract_metadata(text, doc)
        self._segment_chapters(text, doc)
        return doc

    def _extract_metadata(self, text: str, doc: BCRStructuredDocument) -> None:
        m = re.search(r"^(?:Binding Corporate Rules|BCR)[\s\-\n]+(.+)$", text, re.MULTILINE)
        if m: doc.title = m.group(0).strip()[:120]
        m = re.search(r"(?:Version|V)\s*[:\.]?\s*([\d\.]+)", text[:2000], re.IGNORECASE)
        if m: doc.version = m.group(1)
        m = re.search(r"(?:Effective|Effective Date|Date)[\s:]*([A-Z][a-z]+ \d{1,2},? \d{4}|\d{1,2}[\./-]\d{1,2}[\./-]\d{4})", text[:2000])
        if m: doc.effective_date = m.group(1)

    def _segment_chapters(self, text: str, doc: BCRStructuredDocument) -> None:
        parts = self._section_pattern.split(text)
        if len(parts) < 2:
            doc.chapters = [BCRStructuredChapter(chapter_no=1, title="Full Document", content=text)]
            return

        chapter_no = 0
        for part in parts:
            part = part.strip()
            if len(part) < 80:
                continue
            chapter_no += 1
            heading = ""
            lines = part.split("\n")
            if lines:
                heading = lines[0].strip()[:100]
            # Map heading to known section name
            for keyword, canon in self._known_titles.items():
                if keyword in heading.lower():
                    heading = canon.capitalize()
                    break
            doc.chapters.append(BCRStructuredChapter(
                chapter_no=chapter_no, title=heading, content=part,
            ))

        # Detect appendix titles
        for m in re.finditer(r"(?:Annex|Appendix|ANNEX|APPENDIX)\s+([A-Z0-9]+)[\.:\)\s]*([^\n]{0,80})", text):
            doc.appendix_titles.append(m.group(0).strip()[:80])

    @staticmethod
    def _extract_docx(path: Path) -> str:
        return "\n".join(p.text for p in Document(str(path)).paragraphs if p.text.strip())

    @staticmethod
    def _extract_pdf(path: Path) -> str:
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
