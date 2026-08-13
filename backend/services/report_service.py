from pathlib import Path

from docx import Document
from sqlalchemy.orm import Session

from backend.core.json_utils import dumps
from backend.core.settings import Settings
from backend.models.report import ReportArtifactModel
from backend.repositories.report_repository import ReportRepository
from backend.schemas.common import ReportArtifact
from backend.common.render.pdf_renderer import get_pdf_renderer


class ReportService:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings
        self.repository = ReportRepository()

    def create_html_report(
        self,
        db: Session,
        user_id: str,
        owner_type: str,
        owner_id: str,
        filename: str,
        html: str,
        preview: dict,
    ) -> ReportArtifact:
        path = self._owner_dir(owner_type, owner_id) / filename
        path.write_text(html, encoding="utf-8")
        return self._persist(db, user_id, owner_type, owner_id, "html", path, preview)

    def create_markdown_report(
        self,
        db: Session,
        user_id: str,
        owner_type: str,
        owner_id: str,
        filename: str,
        markdown: str,
        preview: dict,
    ) -> ReportArtifact:
        path = self._owner_dir(owner_type, owner_id) / filename
        path.write_text(markdown, encoding="utf-8")
        return self._persist(db, user_id, owner_type, owner_id, "markdown", path, preview)

    def create_pdf_report(
        self,
        db: Session,
        user_id: str,
        owner_type: str,
        owner_id: str,
        filename: str,
        lines: list[str],
        preview: dict,
    ) -> ReportArtifact:
        """Legacy Markdown-line PDF writer (deprecated).

        Kept for limited backward compatibility only. New callers should use
        :meth:`create_pdf_report_ir`, which renders from a v4 ``DocumentIR``.
        """
        path = self._owner_dir(owner_type, owner_id) / filename
        self._write_pdf(path, lines)
        return self._persist(db, user_id, owner_type, owner_id, "pdf", path, preview)

    def create_pdf_report_ir(
        self,
        db: Session,
        user_id: str,
        owner_type: str,
        owner_id: str,
        filename: str,
        document_ir,
        registry,
        preview: dict,
    ) -> ReportArtifact:
        """Render a v4 ``DocumentIR`` to a fixed-layout PDF (task068 T07).

        The renderer consumes the canonical IR directly — never the legacy
        Markdown lines — and degrades the CJK font to ``STSong-Light`` (emb=no)
        until a licensed + hashed CJK font asset is tracked. The PDF font gate
        therefore stays ``BLOCKED_BY_FONT`` on this machine.
        """
        from backend.common.reporting.renderers.pdf import render_pdf

        path = self._owner_dir(owner_type, owner_id) / filename
        path.write_bytes(render_pdf(document_ir, registry))
        return self._persist(db, user_id, owner_type, owner_id, "pdf", path, preview)

    def create_docx_report(
        self,
        db: Session,
        user_id: str,
        owner_type: str,
        owner_id: str,
        filename: str,
        sections: list[tuple[str, list[str]]],
        preview: dict,
    ) -> ReportArtifact:
        """Legacy string-based DOCX writer (deprecated).

        Kept for limited backward compatibility only. New callers should use
        :meth:`create_docx_report_ir`, which renders from a v4 ``DocumentIR``.
        """
        path = self._owner_dir(owner_type, owner_id) / filename
        document = Document()
        for heading, paragraphs in sections:
            document.add_heading(heading, level=1)
            for paragraph in paragraphs:
                document.add_paragraph(paragraph)
        document.save(path)
        return self._persist(db, user_id, owner_type, owner_id, "docx", path, preview)

    def create_docx_report_ir(
        self,
        db: Session,
        user_id: str,
        owner_type: str,
        owner_id: str,
        filename: str,
        document_ir,
        registry,
        preview: dict,
    ) -> ReportArtifact:
        """Render a v4 ``DocumentIR`` to a native DOCX report (task068 T06).

        The renderer consumes the canonical IR directly — never the legacy
        ``sections`` string lists — and produces native Word structures
        (headings/tables/numbering/header/footer/page fields).
        """
        from backend.common.reporting.renderers.docx import render_docx

        path = self._owner_dir(owner_type, owner_id) / filename
        path.write_bytes(render_docx(document_ir, registry))
        return self._persist(db, user_id, owner_type, owner_id, "docx", path, preview)

    def get_owner_artifact(
        self,
        db: Session,
        user_id: str,
        owner_type: str,
        owner_id: str,
        artifact_type: str,
    ) -> ReportArtifact | None:
        model = self.repository.get_by_owner_and_type(db, user_id, owner_type, owner_id, artifact_type)
        return self._to_schema(model) if model else None


    def register_external_artifact(
        self,
        db: Session,
        user_id: str,
        owner_type: str,
        owner_id: str,
        artifact_type: str,
        file_path: str,
        preview: dict,
    ) -> ReportArtifact:
        existing = self.repository.get_by_user_and_path(db, user_id, file_path)
        if existing:
            return self._to_schema(existing)
        model = ReportArtifactModel(
            user_id=user_id,
            owner_type=owner_type,
            owner_id=owner_id,
            artifact_type=artifact_type,
            file_path=file_path,
            preview_json=dumps(preview),
        )
        created = self.repository.create(db, model)
        return self._to_schema(created)

    def _owner_dir(self, owner_type: str, owner_id: str) -> Path:
        path = self.settings.report_dir / owner_type / owner_id
        path.mkdir(parents=True, exist_ok=True)
        return path

    def _write_pdf(self, path: Path, lines: list[str]) -> None:
        get_pdf_renderer().from_markdown("\n\n".join(lines), path, path.stem)

    def _persist(
        self,
        db: Session,
        user_id: str,
        owner_type: str,
        owner_id: str,
        artifact_type: str,
        path: Path,
        preview: dict,
    ) -> ReportArtifact:
        model = ReportArtifactModel(
            user_id=user_id,
            owner_type=owner_type,
            owner_id=owner_id,
            artifact_type=artifact_type,
            file_path=str(path),
            preview_json=dumps(preview),
        )
        created = self.repository.create(db, model)
        return self._to_schema(created)

    def _to_schema(self, model: ReportArtifactModel) -> ReportArtifact:
        from backend.core.json_utils import loads

        return ReportArtifact(
            id=model.id,
            owner_type=model.owner_type,
            owner_id=model.owner_id,
            artifact_type=model.artifact_type,
            file_path=model.file_path,
            preview=loads(model.preview_json, {}),
            created_at=model.created_at,
        )
