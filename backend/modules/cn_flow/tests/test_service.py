import json
from pathlib import Path

from docx import Document

from backend.modules.cn_flow import service as cn_flow_service
from backend.modules.cn_flow.schema import CNFlowRequest
from backend.modules.cn_flow.service import CNFlowService


class _DisabledLLM:
    enabled = False


def test_cn_flow_generate_report() -> None:
    tmp_dir = Path("outputs/cn_flow/_test_templates")
    tmp_dir.mkdir(parents=True, exist_ok=True)
    md_template = tmp_dir / "cn_flow_template.md"
    md_template.write_text("# {{business_overview}}\n\n{{risk_rating}}\n", encoding="utf-8")
    docx_template = tmp_dir / "cn_flow_template.docx"
    document = Document()
    document.add_heading("{{business_overview}}", level=1)
    document.add_paragraph("{{risk_rating}}")
    document.save(docx_template)
    cn_flow_service.TEMPLATE_MD = md_template
    cn_flow_service.TEMPLATE_PATH = docx_template

    service = CNFlowService(llm_client=_DisabledLLM())
    payload = CNFlowRequest.model_validate(
        {
            "company_name": "测试企业",
            "transfer_purpose": "全球客服与风控",
            "data_categories": ["账户信息", "设备信息"],
            "sensitive_data_flags": ["生物识别"],
            "recipient_entities": [
                {
                    "entity_name": "US ServiceCo",
                    "country_region": "United States",
                    "entity_role": "processor",
                    "is_restricted_party": False,
                }
            ],
            "transfer_chain": "CN -> US processor -> subprocessor",
            "attachments": [
                {
                    "file_role": "data_inventory",
                    "file_name": "data.csv",
                    "file_format": "csv",
                    "storage_uri": "storage://uploads/data.csv",
                },
                {
                    "file_role": "entity_inventory",
                    "file_name": "entity.csv",
                    "file_format": "csv",
                    "storage_uri": "storage://uploads/entity.csv",
                },
            ],
        }
    )

    result = service.generate_report(payload)
    assert result.report_path.endswith(".docx")
    assert "_14117_风险评估结论报告_草案_" in result.report_path
    assert result.output_files["pdf"].endswith(".pdf")
    assert result.output_files["xlsx"].endswith(".xlsx")
    assert result.output_files["zip"].endswith(".zip")
    assert result.output_files["facts_json"].endswith(".json")
    assert result.output_files["issue_list_json"].endswith(".json")
    assert result.output_files["evidence_chain_json"].endswith(".json")
    assert result.output_files["trace_manifest"].endswith(".json")
    assert result.risk_items

    manifest = json.loads(Path(result.output_files["trace_manifest"]).read_text(encoding="utf-8"))
    event_names = [item["name"] for item in manifest["events"]]
    assert "cn_flow_request" in event_names
    assert "facts_built" in event_names
    assert "issues_built" in event_names
    assert "evidence_built" in event_names
    assert "context_pack_built" in event_names
