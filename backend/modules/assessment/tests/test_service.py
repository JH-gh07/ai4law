import json
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import openpyxl

from backend.common.knowledge.v2 import RetrievalBundle

from backend.modules.assessment.task_state import AssessmentTaskState
from docx import Document

from backend.modules.assessment import report_renderer
from backend.modules.assessment.retriever import AssessmentRetriever
from backend.modules.assessment.schema import AssessmentRequest
from backend.modules.assessment.service import AssessmentService
from backend.modules.diagnosis.service import DiagnosisService


class _DisabledLLM:
    enabled = False


class _DisabledLegalService:
    enabled = False


def _disable_external_services(monkeypatch) -> None:
    monkeypatch.setattr(DiagnosisService, "_build_rule_explanation", lambda self, result, answers: result.rationale)


def _build_service() -> AssessmentService:
    return AssessmentService(llm_client=_DisabledLLM(), legal_api_service=_DisabledLegalService())


def _install_test_templates(monkeypatch, tmp_path: Path) -> None:
    md_template = tmp_path / "assessment_template.md"
    md_template.write_text(
        "# {{company_name}}\n\n{{business_flow_summary}}\n\n{{overall_conclusion}}\n",
        encoding="utf-8",
    )

    docx_template = tmp_path / "assessment_template.docx"
    document = Document()
    document.add_heading("{{company_name}}", level=0)
    document.add_paragraph("{{business_flow_summary}}")
    document.add_paragraph("{{overall_conclusion}}")
    document.save(docx_template)

    monkeypatch.setattr(report_renderer, "TEMPLATE_MD", md_template)
    monkeypatch.setattr(report_renderer, "TEMPLATE_PATH", docx_template)


def _read_trace_event(manifest_path: str, event_name: str) -> dict:
    manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
    for event in manifest["events"]:
        if event["name"] == event_name:
            return json.loads(Path(event["path"]).read_text(encoding="utf-8"))
    raise AssertionError(f"trace event not found: {event_name}")


def test_assessment_security_assessment_path_generates_report(monkeypatch, tmp_path) -> None:
    _disable_external_services(monkeypatch)
    _install_test_templates(monkeypatch, tmp_path)
    service = _build_service()
    payload = AssessmentRequest(
        company_name="测试公司",
        industry="医疗科技",
        is_ciio=True,
        contains_important_data=False,
        pii_count=200000,
        spi_count=300,
        transfer_purpose="跨境客服",
        receiver_country="Singapore",
        force_override_path=False,
        uploaded_files=[],
    )

    result = service.generate_report(payload)

    assert result.state == "COMPLETED"
    assert len(result.chapters) == 8
    assert result.report_path.endswith(".docx")
    assert "_数据出境风险自评估报告_草案_" in result.report_path
    assert result.output_files["markdown"].endswith(".md")
    assert "_数据出境风险自评估报告_草案_" in result.output_files["markdown"]
    assert result.output_files["zip"].endswith(".zip")
    assert Path(result.output_files["markdown"]).exists()
    assert Path(result.output_files["docx"]).exists()
    assert Path(result.output_files["zip"]).exists()
    assert Path(result.output_files["trace_manifest"]).exists()

    with ZipFile(result.output_files["zip"]) as bundle:
        names = set(bundle.namelist())
    assert Path(result.output_files["markdown"]).name in names
    assert Path(result.output_files["docx"]).name in names

    # Phase 7: intermediate artifacts
    assert result.output_files["issue_list_json"].endswith(".json")
    assert result.output_files["issue_list_xlsx"].endswith(".xlsx")
    assert result.output_files["evidence_chain_json"].endswith(".json")
    assert result.output_files["evidence_chain_xlsx"].endswith(".xlsx")
    assert Path(result.output_files["issue_list_json"]).exists()
    assert Path(result.output_files["issue_list_xlsx"]).exists()
    assert Path(result.output_files["evidence_chain_json"]).exists()
    assert Path(result.output_files["evidence_chain_xlsx"]).exists()
    assert "trace_manifest" in result.output_files
    assert Path(result.output_files["trace_manifest"]).exists()

    # XLSX headers
    issue_wb = openpyxl.load_workbook(result.output_files["issue_list_xlsx"])
    assert list(next(issue_wb.active.iter_rows(min_row=1, max_row=1, values_only=True))) == [
        "问题编号", "标题", "描述", "类别", "严重程度", "事实引用", "规则引用", "证据引用", "建议措施", "影响输出",
    ]
    evidence_wb = openpyxl.load_workbook(result.output_files["evidence_chain_xlsx"])
    assert list(next(evidence_wb.active.iter_rows(min_row=1, max_row=1, values_only=True))) == [
        "证据编号", "主张", "事实引用", "规则引用", "结论", "置信度", "被使用于",
    ]

    # All artifacts in ZIP
    for key in ("issue_list_json", "issue_list_xlsx", "evidence_chain_json", "evidence_chain_xlsx", "trace_manifest"):
        assert Path(result.output_files[key]).name in names

    # Phase 8: facts, path judgment, material checklist JSON
    assert result.output_files["facts_json"].endswith(".json")
    assert result.output_files["path_judgment_json"].endswith(".json")
    assert result.output_files["material_checklist_json"].endswith(".json")
    assert result.output_files["material_checklist_xlsx"].endswith(".xlsx")
    assert Path(result.output_files["facts_json"]).exists()
    assert Path(result.output_files["path_judgment_json"]).exists()
    assert Path(result.output_files["material_checklist_json"]).exists()
    assert Path(result.output_files["material_checklist_xlsx"]).exists()
    for key in ("facts_json", "path_judgment_json", "material_checklist_json", "material_checklist_xlsx"):
        assert Path(result.output_files[key]).name in names

    import json as _json
    facts_data = _json.loads(Path(result.output_files["facts_json"]).read_text(encoding="utf-8"))
    assert len(facts_data) > 0
    assert "fact_id" in facts_data[0]
    path_data = _json.loads(Path(result.output_files["path_judgment_json"]).read_text(encoding="utf-8"))
    assert "recommended_path" in path_data
    assert "risk_level" in path_data
    assert path_data["is_override"] is False
    materials_data = _json.loads(Path(result.output_files["material_checklist_json"]).read_text(encoding="utf-8"))
    assert materials_data
    assert materials_data[0]["status"] == "待补充"

    request_event = _read_trace_event(result.output_files["trace_manifest"], "assessment_request")
    assert request_event["payload"]["company_name"] == "测试公司"
    facts_event = _read_trace_event(result.output_files["trace_manifest"], "facts_built")
    assert facts_event["payload"]["facts"]
    issues_event = _read_trace_event(result.output_files["trace_manifest"], "issues_built")
    assert issues_event["payload"]["issues"]
    evidence_event = _read_trace_event(result.output_files["trace_manifest"], "evidence_built")
    assert evidence_event["payload"]["evidence_chain"]
    evidence_ids = {item["evidence_id"] for item in evidence_event["payload"]["evidence_chain"]}
    for issue in evidence_event["payload"]["issues"]:
        assert set(issue["evidence_refs"]).issubset(evidence_ids)
    context_event = _read_trace_event(result.output_files["trace_manifest"], "context_pack_built")
    assert context_event["payload"]["module_key"] == "assessment"
    assert context_event["payload"]["evidence_chain"]
    consistency_event = _read_trace_event(result.output_files["trace_manifest"], "consistency_issues")
    assert "context_pack_references" in consistency_event["payload"]["checks"]
    assert "report_against_context" in consistency_event["payload"]["checks"]



def test_assessment_does_not_reject_path(monkeypatch, tmp_path) -> None:
    """Path judgment has been moved upstream to the diagnosis module.
    Assessment module now always generates the report regardless of path."""
    _disable_external_services(monkeypatch)
    _install_test_templates(monkeypatch, tmp_path)
    service = _build_service()
    payload = AssessmentRequest(
        company_name="测试公司",
        industry="SaaS",
        is_ciio=False,
        contains_important_data=False,
        pii_count=200000,
        spi_count=300,
        transfer_purpose="跨境客服",
        receiver_country="Singapore",
        force_override_path=False,
        uploaded_files=[],
        path_check_mode="generate_only",  # default — always generate
    )

    result = service.generate_report(payload)
    # Should succeed — path judgment is upstream
    assert result.state == AssessmentTaskState.COMPLETED
    assert result.report_path


def test_assessment_generates_regardless_of_path(monkeypatch, tmp_path) -> None:
    """Assessment module always generates the report.
    Path consistency is recorded in trace but does not block generation."""
    _disable_external_services(monkeypatch)
    _install_test_templates(monkeypatch, tmp_path)
    service = _build_service()
    payload = AssessmentRequest(
        company_name="测试公司",
        industry="SaaS",
        is_ciio=False,
        contains_important_data=False,
        pii_count=200000,
        spi_count=300,
        transfer_purpose="跨境客服",
        receiver_country="Singapore",
        force_override_path=True,
        uploaded_files=[],
    )

    result = service.generate_report(payload)
    # Should succeed — assessment module always generates
    assert result.state == AssessmentTaskState.COMPLETED
    assert result.report_path


def test_assessment_zip_has_no_duplicate_names(monkeypatch, tmp_path) -> None:
    _disable_external_services(monkeypatch)
    _install_test_templates(monkeypatch, tmp_path)
    service = _build_service()
    payload = AssessmentRequest(
        company_name="测试公司",
        industry="医疗科技",
        is_ciio=True,
        contains_important_data=False,
        pii_count=200000,
        spi_count=300,
        transfer_purpose="跨境客服",
        receiver_country="Singapore",
        force_override_path=False,
        uploaded_files=[],
    )
    result = service.generate_report(payload)
    with ZipFile(result.output_files["zip"]) as bundle:
        names = bundle.namelist()
    assert len(names) == len(set(names))


def test_assessment_retriever_query_contains_profile_fields() -> None:
    captured: dict[str, object] = {}

    class StubRetrievalService:
        def retrieve_with_fallback(self, request, **kwargs):
            captured["request"] = request
            captured["kwargs"] = kwargs
            return SimpleNamespace(
                bundle=RetrievalBundle(),
                manifest=None,
            )

    profile = _build_service().extractor.extract(
        AssessmentRequest(
            company_name="测试公司",
            industry="互联网医疗",
            is_ciio=True,
            contains_important_data=True,
            pii_count=100,
            spi_count=50,
            transfer_purpose="模型训练",
            receiver_country="新加坡",
        )
    )

    hits = AssessmentRetriever(
        retrieval_service=StubRetrievalService(),
    ).search(profile)

    assert hits == []
    request = captured["request"]
    assert "互联网医疗" in request.query
    assert "模型训练" in request.query
    assert "新加坡" in request.query
    assert "CIIO" in request.query
    assert "important data" in request.query
    assert request.jurisdiction == "cn"
    assert request.path == "assessment"


def test_assessment_fails_when_issue_builder_is_unavailable(monkeypatch, tmp_path) -> None:
    _disable_external_services(monkeypatch)
    _install_test_templates(monkeypatch, tmp_path)
    service = _build_service()

    def _broken_issue_builder(*args, **kwargs):
        raise RuntimeError("issue builder unavailable")

    monkeypatch.setattr("backend.modules.assessment.service.build_assessment_issues", _broken_issue_builder)

    payload = AssessmentRequest(
        company_name="测试公司",
        industry="医疗科技",
        is_ciio=True,
        contains_important_data=False,
        pii_count=200000,
        spi_count=300,
        transfer_purpose="跨境客服",
        receiver_country="Singapore",
        force_override_path=False,
        uploaded_files=[],
    )

    try:
        service.generate_report(payload)
        raise AssertionError("expected RuntimeError when issue builder is unavailable")
    except RuntimeError as exc:
        assert "issue builder unavailable" in str(exc)
