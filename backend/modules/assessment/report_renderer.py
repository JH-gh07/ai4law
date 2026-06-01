from __future__ import annotations

from pathlib import Path
from shutil import copy2
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl

from backend.common.render.report import (
    format_date_stamp,
    render_docx_template,
    render_markdown_template,
    safe_filename,
)
from backend.common.render.summary import attach_citations, dedup_if_same, summarize_for_slot
from typing import TYPE_CHECKING, Any

from backend.common.workflow.evidence import EvidenceItem
from backend.common.workflow.facts import FactItem
from backend.common.workflow.issues import IssueItem
from backend.modules.assessment.internal_review_generator import (
    build_internal_review_payload,
    generate_internal_review_markdown,
)
from backend.modules.assessment.schema import ChapterContent, CompanyProfile, RegulationHit

if TYPE_CHECKING:
    from backend.common.citation.registry import CitationRegistry

TEMPLATE_PATH = Path("doc/v2/assets/templates/2.2_risk_assessment_template_v0.docx")
TEMPLATE_MD = Path("doc/v2/assets/templates/2.2_risk_assessment_template_v0.md")

_ISSUE_XLSX_HEADERS = [
    "问题编号",
    "标题",
    "描述",
    "类别",
    "严重程度",
    "事实引用",
    "规则引用",
    "证据引用",
    "建议措施",
    "影响输出",
]

_EVIDENCE_XLSX_HEADERS = [
    "证据编号",
    "主张",
    "事实引用",
    "规则引用",
    "结论",
    "置信度",
    "被使用于",
]

_MATERIAL_XLSX_HEADERS = ["材料编号", "摘要", "状态"]


class AssessmentReportRenderer:

    def render(
        self,
        task_id: str,
        company_name: str,
        profile: CompanyProfile,
        regulations: list[RegulationHit],
        chapters: list[ChapterContent],
        path_warning: str | None = None,
        alignment_warning: str | None = None,
        issues: list[IssueItem] | None = None,
        evidence_chain: list[EvidenceItem] | None = None,
        attachment_notes: list[dict[str, str]] | None = None,
        trace_manifest_path: str | None = None,
        facts: list[FactItem] | None = None,
        diagnosis_result: dict | None = None,
        force_override: bool = False,
        writing_strategy: dict[str, Any] | None = None,
        generation_basis_pack: dict[str, Any] | None = None,
        legal_grounding: dict[str, Any] | None = None,
        citation_registry: "CitationRegistry | None" = None,
    ) -> dict[str, str]:
        output_dir = Path("outputs/assessment") / task_id / "outputs"
        output_dir.mkdir(parents=True, exist_ok=True)
        date_stamp = format_date_stamp()
        safe_company = safe_filename(company_name)
        md_output = output_dir / f"{safe_company}_数据出境风险自评估报告_草案_{date_stamp}.md"
        docx_output = output_dir / f"{safe_company}_数据出境风险自评估报告_草案_{date_stamp}.docx"
        zip_output = output_dir / f"{safe_company}_安全评估路径输出包_草案_{date_stamp}.zip"
        mapping = _build_template_mapping(
            profile, regulations, chapters, date_stamp, path_warning, alignment_warning,
            citation_registry=citation_registry,
        )
        render_markdown_template(md_output, TEMPLATE_MD, mapping)
        render_docx_template(docx_output, TEMPLATE_PATH, mapping)

        result: dict[str, str] = {
            "markdown": str(md_output),
            "docx": str(docx_output),
            "zip": str(zip_output),
        }

        if issues:
            result["issue_list_json"] = _write_issue_list_json(issues, output_dir)
            result["issue_list_xlsx"] = _write_issue_list_xlsx(issues, output_dir)

        if evidence_chain:
            result["evidence_chain_json"] = _write_evidence_chain_json(evidence_chain, output_dir)
            result["evidence_chain_xlsx"] = _write_evidence_chain_xlsx(evidence_chain, output_dir)

        material_rows = _build_material_checklist_rows(
            attachment_notes or [],
            issues or [],
            facts or [],
        )
        result["material_checklist_xlsx"] = _write_material_checklist_xlsx(material_rows, output_dir)
        result["material_checklist_json"] = _write_material_checklist_json(material_rows, output_dir)

        if facts:
            result["facts_json"] = _write_facts_json(facts, output_dir)

        if diagnosis_result:
            result["path_judgment_json"] = _write_path_judgment_json(
                diagnosis_result,
                path_warning,
                output_dir,
                force_override=force_override,
            )

        if trace_manifest_path:
            dest = output_dir / "trace_manifest.json"
            copy2(trace_manifest_path, dest)
            result["trace_manifest"] = str(dest)

        # Enhanced outputs: internal review, legal grounding, writing strategy, generation basis pack
        if issues and writing_strategy and generation_basis_pack:
            internal_review_md = _write_internal_review(
                issues=issues,
                writing_strategy=writing_strategy,
                generation_basis_pack=generation_basis_pack,
                material_rows=material_rows,
                output_dir=output_dir,
            )
            result["internal_review_md"] = str(internal_review_md)

        if legal_grounding:
            result["legal_grounding_json"] = _write_legal_grounding_json(legal_grounding, output_dir)

        if writing_strategy:
            result["writing_strategy_json"] = _write_writing_strategy_json(writing_strategy, output_dir)

        if generation_basis_pack:
            result["generation_basis_pack_json"] = _write_generation_basis_pack_json(generation_basis_pack, output_dir)

        if citation_registry is not None:
            result["citation_map_json"] = _write_citation_map_json(citation_registry, output_dir)

        with ZipFile(zip_output, mode="w", compression=ZIP_DEFLATED) as bundle:
            for key, path in result.items():
                if key == "zip":
                    continue
                p = Path(path)
                if p.exists():
                    bundle.write(p, arcname=p.name)

        return result


def _build_template_mapping(
    profile: CompanyProfile,
    regulations: list[RegulationHit],
    chapters: list[ChapterContent],
    date_stamp: str,
    path_warning: str | None = None,
    alignment_warning: str | None = None,
    citation_registry: "CitationRegistry | None" = None,
) -> dict[str, str]:
    def pick(title: str) -> str:
        for chapter in chapters:
            if chapter.title == title:
                return chapter.content
        return ""

    pii = f"{profile.pii_count:,}人"
    spi = f"{profile.spi_count:,}人"
    contains_spi = "是" if profile.spi_count > 0 else "否"

    legal_basis = summarize_for_slot(pick("出境必要性与合法性基础"), max_sentences=2, max_chars=260)
    security_measures = summarize_for_slot(pick("安全措施与传输机制"), max_sentences=2, max_chars=260)
    risk_summary = summarize_for_slot(pick("剩余风险与整改建议"), max_sentences=2, max_chars=260)
    conclusion = summarize_for_slot(pick("综合评估结论"), max_sentences=2, max_chars=260)
    if path_warning:
        conclusion = f"【路径不匹配警示】{path_warning}\n{conclusion}".strip()
    if alignment_warning:
        conclusion = f"【输入对齐告警】{alignment_warning}\n{conclusion}".strip()
    citations = [f"{item.title}{item.article}" for item in regulations[:3]]
    regulations_text = "; ".join(citations) or "未提供"

    governance_candidate = summarize_for_slot(pick("剩余风险与整改建议"), max_sentences=1, max_chars=200)
    technical_candidate = summarize_for_slot(pick("安全措施与传输机制"), max_sentences=1, max_chars=200)
    governance_measures, technical_measures = dedup_if_same(
        governance_candidate,
        technical_candidate,
        "管理措施待补充（当前仅提供技术措施摘要）。",
    )

    business_flow_summary = summarize_for_slot(pick("出境活动概述"), max_sentences=2, max_chars=260) or "基于企业业务需要进行跨境传输。"
    legal_risk = legal_basis or risk_summary
    security_risk = security_measures or "需结合技术与管理措施持续评估。"
    process_risk = risk_summary or "流程风险可控，但需完善内控。"
    governance_measures = attach_citations(governance_measures, citations)
    technical_measures = attach_citations(technical_measures, citations)
    business_flow_summary = attach_citations(business_flow_summary, citations)
    legal_risk = attach_citations(legal_risk, citations)
    security_risk = attach_citations(security_risk, citations)
    process_risk = attach_citations(process_risk, citations)
    conclusion = attach_citations(conclusion or "综合评估结论需结合监管要求进一步确认。", citations)

    return {
        "company_name": profile.company_name,
        "company_uscc": "未提供",
        "report_date": date_stamp,
        "contact_name": "未提供",
        "transfer_purpose": profile.transfer_purpose,
        "recipient_name": profile.receiver_country or "未提供",
        "recipient_country_region": profile.receiver_country or "未提供",
        "business_flow_summary": business_flow_summary,
        "data_categories": f"个人信息（普通/敏感），参考法规：{regulations_text}",
        "data_volume": f"普通个人信息{pii}，敏感个人信息{spi}",
        "contains_spi": contains_spi,
        "transfer_frequency": "未提供",
        "legal_risk": legal_risk,
        "security_risk": security_risk,
        "process_risk": process_risk,
        "governance_measures": governance_measures,
        "technical_measures": technical_measures,
        "contractual_measures": "拟通过合同与附加条款约束接收方处理范围与责任。",
        "overall_conclusion": conclusion,
        "remediation_items": attach_citations("详见风险识别与整改建议章节。", citations),
        "attachments": "- 无",
        "citation_map": citation_registry.build_citation_map_section() if citation_registry else "",
    }


# ---------------------------------------------------------------------------
# Phase 7: intermediate artifact writers
# ---------------------------------------------------------------------------


def _write_issue_list_json(issues: list[IssueItem], output_dir: Path) -> str:
    import json

    path = output_dir / "issue_list.json"
    path.write_text(json.dumps([issue.model_dump() for issue in issues], ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def _write_issue_list_xlsx(issues: list[IssueItem], output_dir: Path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "问题清单"
    ws.append(_ISSUE_XLSX_HEADERS)
    for issue in issues:
        ws.append([
            issue.issue_id,
            issue.title,
            issue.description,
            issue.category,
            issue.severity,
            "; ".join(issue.fact_refs),
            "; ".join(issue.rule_refs),
            "; ".join(issue.evidence_refs),
            issue.recommended_action,
            "; ".join(issue.affects_outputs),
        ])
    path = output_dir / "issue_list.xlsx"
    wb.save(path)
    return str(path)


def _write_evidence_chain_json(evidence_chain: list[EvidenceItem], output_dir: Path) -> str:
    import json

    path = output_dir / "evidence_chain.json"
    path.write_text(
        json.dumps([ev.model_dump() for ev in evidence_chain], ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return str(path)


def _write_evidence_chain_xlsx(evidence_chain: list[EvidenceItem], output_dir: Path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "证据链"
    ws.append(_EVIDENCE_XLSX_HEADERS)
    for ev in evidence_chain:
        ws.append([
            ev.evidence_id,
            ev.claim,
            "; ".join(ev.fact_refs),
            "; ".join(ev.rule_refs),
            ev.conclusion,
            ev.confidence,
            "; ".join(ev.used_by),
        ])
    path = output_dir / "evidence_chain.xlsx"
    wb.save(path)
    return str(path)


def _write_material_checklist_xlsx(material_rows: list[dict[str, str]], output_dir: Path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "材料补充清单"
    ws.append(_MATERIAL_XLSX_HEADERS)
    for note in material_rows:
        ws.append([
            note.get("source_ref", ""),
            note.get("summary", ""),
            note.get("status", "待补充"),
        ])
    path = output_dir / "material_checklist.xlsx"
    wb.save(path)
    return str(path)


# ---------------------------------------------------------------------------
# Phase 8: additional intermediate artifact writers
# ---------------------------------------------------------------------------


def _write_facts_json(facts: list[FactItem], output_dir: Path) -> str:
    import json

    path = output_dir / "facts.json"
    path.write_text(
        json.dumps([fact.model_dump() for fact in facts], ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return str(path)


def _write_path_judgment_json(
    diagnosis_result: dict,
    path_warning: str | None,
    output_dir: Path,
    force_override: bool = False,
) -> str:
    import json

    payload = {
        "recommended_path": diagnosis_result.get("recommended_path"),
        "risk_level": diagnosis_result.get("risk_level"),
        "rationale": diagnosis_result.get("rationale"),
        "path_warning": path_warning,
        "is_override": force_override,
    }
    path = output_dir / "path_judgment.json"
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def _write_material_checklist_json(material_rows: list[dict[str, str]], output_dir: Path) -> str:
    import json

    path = output_dir / "material_checklist.json"
    path.write_text(json.dumps(material_rows, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def _build_material_checklist_rows(
    attachment_notes: list[dict[str, str]],
    issues: list[IssueItem],
    facts: list[FactItem],
) -> list[dict[str, str]]:
    if attachment_notes:
        return [
            {
                "source_ref": note.get("source_ref", ""),
                "summary": note.get("summary", ""),
                "status": "待补充",
            }
            for note in attachment_notes
        ]

    uploaded_files_fact = next(
        (fact for fact in facts if fact.field_path == "request.uploaded_files"),
        None,
    )
    uploaded_files = uploaded_files_fact.normalized_value if uploaded_files_fact else []
    if isinstance(uploaded_files, list) and uploaded_files:
        return [
            {
                "source_ref": str(item),
                "summary": "已提供附件路径，但尚未形成附件解析摘要，需补充材料审查结果。",
                "status": "待解析",
            }
            for item in uploaded_files
        ]

    missing_attachments_issue = next(
        (issue for issue in issues if issue.issue_id == "ISSUE-missing-attachments"),
        None,
    )
    if missing_attachments_issue:
        return [
            {
                "source_ref": "uploaded_files",
                "summary": missing_attachments_issue.recommended_action,
                "status": "待补充",
            }
        ]

    return [
        {
            "source_ref": "attachment_summary",
            "summary": "当前未形成附件解析摘要；如涉及申报材料，请补充数据清单、隐私政策、合同/协议与安全措施说明。",
            "status": "待补充",
        }
    ]


# ---------------------------------------------------------------------------
# Batch 3: enhanced output artifact writers
# ---------------------------------------------------------------------------


def _write_internal_review(
    *,
    issues: list[IssueItem],
    writing_strategy: dict[str, Any],
    generation_basis_pack: dict[str, Any],
    material_rows: list[dict[str, str]],
    output_dir: Path,
) -> Path:
    payload = build_internal_review_payload(
        issues=issues,
        writing_strategy=writing_strategy,
        generation_basis_pack=generation_basis_pack,
        material_rows=material_rows,
    )
    content = generate_internal_review_markdown(payload=payload, llm_client=None)
    path = output_dir / "internal_ai_review.md"
    path.write_text(content, encoding="utf-8")
    return path


def _write_legal_grounding_json(legal_grounding: dict[str, Any], output_dir: Path) -> str:
    import json

    path = output_dir / "legal_grounding.json"
    path.write_text(json.dumps(legal_grounding, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def _write_writing_strategy_json(writing_strategy: dict[str, Any], output_dir: Path) -> str:
    import json

    path = output_dir / "writing_strategy.json"
    path.write_text(json.dumps(writing_strategy, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def _write_generation_basis_pack_json(generation_basis_pack: dict[str, Any], output_dir: Path) -> str:
    import json

    path = output_dir / "generation_basis_pack.json"
    path.write_text(json.dumps(generation_basis_pack, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def _write_citation_map_json(citation_registry: "CitationRegistry", output_dir: Path) -> str:
    import json

    footnote_map = citation_registry.get_footnote_map()
    payload = {
        "footnote_map": {
            str(num): item.to_dict() for num, item in footnote_map.items()
        },
        "all_items": citation_registry.to_list(),
    }
    path = output_dir / "citation_map.json"
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)
